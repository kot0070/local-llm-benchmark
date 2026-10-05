"""W_GPU_UUID: fingerprint query drops uuid; old manifests still accepted.

All subprocess interaction is mocked (bench.env._run); no real nvidia-smi,
wmic, powershell, or Ollama is ever invoked.
"""
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from bench import env as envmod

UUID_PAT = re.compile(r"GPU-[0-9a-fA-F]{8}-")

OLD_GPU = ("NVIDIA GeForce RTX 3070, GPU-12345678-1234-1234-1234-123456789abc,"
           " 566.92, 8192 MiB, 220.00 W")
NEW_GPU = "NVIDIA GeForce RTX 3070, 566.92, 8192 MiB, 220.00 W"


def _fp_with_gpu(monkeypatch, gpu_line):
    seen = {}

    def fake(cmd, timeout=30.0):
        text = " ".join(cmd)
        if "nvidia-smi" in text and "query-gpu" in text:
            seen["q"] = text
            return True, gpu_line + "\n"
        return False, "mocked-off"
    monkeypatch.setattr(envmod, "_run", fake)
    return envmod.fingerprint(None), seen


def test_fingerprint_query_omits_uuid(monkeypatch):
    fp, seen = _fp_with_gpu(monkeypatch, NEW_GPU)
    assert "uuid" not in seen["q"].lower()
    assert "name,driver_version,memory.total,power.limit" in seen["q"]
    assert fp["gpu"] == NEW_GPU


def test_fingerprint_gpu_has_no_uuid(monkeypatch):
    fp, _ = _fp_with_gpu(monkeypatch, NEW_GPU)
    assert not UUID_PAT.search(fp["gpu"])


def test_owner_summary_accepts_old_and_new():
    from tools import owner_summary as om
    for gpu in (OLD_GPU, NEW_GPU):
        man = {"run_id": "r", "fingerprint": {"gpu": gpu, "cpu": "c",
                                              "ram": "8 GB", "os": "o"}}
        lines = om._overview_lines("r", man, [])
        assert any("GPU:" in ln for ln in lines)


def test_excel_overviews_accept_old_and_new():
    import openpyxl
    from tools import build_excel_report as bu
    from tools import build_excel_report_en as be
    for gpu in (OLD_GPU, NEW_GPU):
        man = {"run_id": "r", "start_local": "", "budget_hours": 8,
               "ollama_version": "x",
               "fingerprint": {"gpu": gpu, "cpu": "c", "ram": "8 GB"}}
        wb = openpyxl.Workbook()
        bu._sheet_overview(wb, "r", man, [], [])
        wb2 = openpyxl.Workbook()
        be._sheet_overview(wb2, "r", man, [], [])
        for _wb in (wb, wb2):
            for ws in _wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str):
                            assert not UUID_PAT.search(v), (ws.title, v)


def test_runner_manifest_stores_fingerprint_opaquely(tmp_path):
    from bench import runner as rmod

    class _Args:
        run_id = "uuidtest"
        budget_hours = 8
        smoke = False

    class _Ctx:
        digest_by_tag = {}
        ollama_version = "x"
        load_info = {}
        size_by_tag = {}

    for gpu in (OLD_GPU, NEW_GPU):
        run_dir = str(tmp_path / ("u" + str(abs(hash(gpu)) % 10**8)))
        ctx = _Ctx()
        ctx.run_dir = run_dir
        os.makedirs(run_dir, exist_ok=True)
        man = rmod.write_run_manifest(ctx, _Args(), ROOT, [], {"gpu": gpu},
                                      {"baseline": {}})
        assert man["fingerprint"]["gpu"] == gpu
        rmod.verify_against_manifest(ctx, ROOT)
