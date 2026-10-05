"""Tests for the 28-model English workbook (W_EN28). No network, no Ollama."""
import glob
import hashlib
import inspect
import os
import re
import subprocess
import sys

import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from tools.build_excel_report_en import (  # noqa: E402
    ELIGIBILITY_N2_RELPATH,
    NIGHT2_TAG_PREFIX,
    SHEET_NAMES,
    TRUNCATION_MARKER,
    UNVERIFIABLE_QUOTE,
    XLSX_NAME_28,
    _load_records_scoped,
    british_hits,
    build_workbook,
    normalize_evidence_quote,
    strip_trailing_marker,
)
from tools.en_sheets_detail import (  # noqa: E402
    DETAIL_HEADER_28,
    FAILURE_HEADER_28,
)
from tools.en_sheets_narrative import (  # noqa: E402
    FORMULA_N1,
    FORMULA_N2,
    GROUP_FILES,
    GROUP_FILES_ALL,
    GROUP_N2_FILENAME,
    N2_BAND_LABEL,
    TIER_FOCUSED,
    TIER_GENERALIST,
    TIER_SINGLE,
    build_all_narrative_28,
    load_analysis,
    load_analysis_28,
)

BASE_RUN = os.path.join(ROOT, "results", "night_20260921-155146")
N2_RUN = os.path.join(ROOT, "results", "night2_b")
BASE_ID = "night_20260921-155146"
N2_ID = "night2_b"
OUT = os.path.join(N2_RUN, XLSX_NAME_28)
NIGHT1_XLSX = os.path.join(BASE_RUN, "NIGHT1_REPORT_EN.xlsx")

EXPECTED_N2_TAGS = frozenset((
    "night2-ministral-3-8b:q4km",
    "night2-nomic-embed-code:q4km",
    "night2-qwen3-vl-8b:q4km",
    "night2-qwen3.5-9b:q4km",
))

# A7 baseline: captured at import, before any test in this module builds.
with open(NIGHT1_XLSX, "rb") as _f:
    NIGHT1_SHA_AT_IMPORT = hashlib.sha256(_f.read()).hexdigest()


def _sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def _open_xlsx(path):
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True)


def _build28():
    return build_workbook(BASE_RUN, OUT, N2_RUN)


@pytest.fixture(scope="module")
def built():
    _build28()
    return OUT


def _scorecard_parts(ws):
    """Return (bands, data_rows) in sheet order."""
    bands = []
    data = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        if all(v is None for v in row[1:]):
            bands.append(row[0])
        else:
            data.append(list(row))
    return bands, data


def test_a1_exists_sheets_order_deterministic(built):
    assert os.path.basename(built) == XLSX_NAME_28
    assert os.path.dirname(os.path.abspath(built)) == os.path.abspath(N2_RUN)
    first = _sha(built)
    _build28()
    second = _sha(built)
    assert first == second
    print("NIGHT28 sha256: %s" % first)
    wb = _open_xlsx(built)
    assert wb.sheetnames == SHEET_NAMES
    assert len(SHEET_NAMES) == len(wb.sheetnames)
    leftovers = glob.glob(os.path.join(N2_RUN, "NIGHT1_REPORT*.xlsx"))
    assert leftovers == [], leftovers


def test_a2_scorecard_bands_formulas_sort(built):
    wb = _open_xlsx(built)
    ws = wb["Model Scorecard"]
    header = [c.value for c in ws[3]]
    assert header[-1] == "Run"
    bands, data = _scorecard_parts(ws)
    assert bands == [TIER_GENERALIST, TIER_FOCUSED, TIER_SINGLE,
                     N2_BAND_LABEL]
    assert len(data) == 28
    assert len({r[0] for r in data}) == 28
    base_rows = [r for r in data if r[-1] == BASE_ID]
    n2_rows = [r for r in data if r[-1] == N2_ID]
    assert len(base_rows) + len(n2_rows) == len(data)
    assert {r[0] for r in n2_rows} == set(EXPECTED_N2_TAGS)
    assert not ({r[0] for r in base_rows} & set(EXPECTED_N2_TAGS))
    # No cross-formula sort: every NIGHT-2 row sits below every NIGHT-1 row.
    order = [r[-1] for r in data]
    assert order == [BASE_ID] * len(base_rows) + [N2_ID] * len(n2_rows)
    # Rating sorts descending only within one tier/band.
    current = None
    last = None
    for row, band in _walk_scorecard(ws):
        if band is not None:
            current = band
            last = None
            continue
        assert current is not None
        if last is not None:
            assert row[1] <= last, (row[0], row[1], last)
        last = row[1]
    note = ws["A2"].value or ""
    assert FORMULA_N1 in note
    assert FORMULA_N2 in note


def _walk_scorecard(ws):
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        if all(v is None for v in row[1:]):
            yield list(row), row[0]
        else:
            yield list(row), None


def test_a3_detail_rows_runs_overview(built):
    wb = _open_xlsx(built)
    ws = wb["Per-Model Test Detail"]
    assert list(ws[1][i].value for i in range(len(DETAIL_HEADER_28))) == \
        DETAIL_HEADER_28
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 672
    assert len({(r[0], r[1]) for r in rows}) == 672
    runs = {r[7] for r in rows}
    assert runs == {BASE_ID, N2_ID}
    assert all(r[7] for r in rows)
    nas = [r for r in rows if r[5] == "Not attempted"]
    assert len(nas) == 394  # 361 base not-attempted + 33 NIGHT-2
    for r in nas:
        assert r[3] is None, r[:2]
        assert r[4] is None, r[:2]
    texts = " ".join(str(c.value or "") for wsn in ("Run Overview",)
                     for row in wb[wsn].iter_rows() for c in row)
    assert BASE_ID in texts
    assert N2_ID in texts
    assert "0.34.2" in texts
    assert "0.34.3" in texts
    assert "720" in texts
    ov = {r[0]: r[1] for r in wb["Run Overview"].iter_rows(values_only=True)
          if r[0]}
    assert ov["Out-of-scope rows excluded"] == 720
    assert ov["run_id (NIGHT-1)"] == BASE_ID
    assert ov["run_id (NIGHT-2)"] == N2_ID


def test_a4_failure_quotes_verified(built):
    from bench import report as rep
    wb = _open_xlsx(built)
    ws = wb["Failure Analysis"]
    assert list(ws[1][i].value for i in range(len(FAILURE_HEADER_28))) == \
        FAILURE_HEADER_28
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len({r[0] for r in rows}) == 28
    index = {}
    for run in (BASE_RUN, N2_RUN):
        for r in rep.load_records(run):
            tag = rep.model_of(r)
            parts = str(r.get("key", "")).split("|")
            case = parts[4] if len(parts) >= 5 else ""
            index.setdefault((tag, case),
                             (r.get("response") or {}).get("content", ""))
    quoted = complete = truncated = 0
    for r in rows:
        q = r[6]
        if q in (None, "Not available") or q in (
                "No failure pattern reached the reporting threshold "
                "in this run",
                "No output \u2014 the model produced no content for this case"):
            continue
        quoted += 1
        assert q != UNVERIFIABLE_QUOTE, (r[0], r[5])
        assert "[...]" not in q, (r[0], r[5])
        stripped = strip_trailing_marker(q)
        assert "[truncated]" not in stripped, (r[0], r[5])
        full = index.get((r[0], r[5]))
        assert full is not None, (r[0], r[5])
        assert stripped, (r[0], r[5])
        if stripped == full:
            complete += 1
        else:
            assert full.startswith(stripped), (r[0], r[5])
            assert q.endswith(TRUNCATION_MARKER), (r[0], r[5])
            truncated += 1
    assert quoted >= 1
    assert complete + truncated == quoted


def test_a5_no_cyrillic_no_british_outside_quotes(built):
    import re
    cyr = re.compile(r"[\u0400-\u04FF]")
    wb = _open_xlsx(built)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if ws.title == "Failure Analysis" and cell.column == 7:
                    continue
                v = cell.value
                if isinstance(v, str) and v:
                    assert not cyr.search(v), (ws.title, cell.coordinate)
                    assert british_hits(v) == [], (
                        ws.title, cell.coordinate, v[:120])


def test_a6_suite_collect_count():
    proc = subprocess.run(
        [sys.executable, "-m", "pytest", "tests/", "--collect-only", "-q"],
        capture_output=True, text=True, cwd=ROOT, timeout=120)
    assert proc.returncode == 0, proc.stderr[-2000:]
    found = re.findall(r"(\d+)\s+tests?\s+collected", proc.stdout)
    assert found, proc.stdout[-2000:]
    assert int(found[-1]) >= 387


def test_a7_night1_untouched_and_backward_compat(built):
    assert _sha(NIGHT1_XLSX) == NIGHT1_SHA_AT_IMPORT
    sig = inspect.signature(build_workbook)
    params = list(sig.parameters.values())
    assert [p.name for p in params][:2] == ["run_dir", "out_path"]
    assert params[2].name == "extra_run_dir"
    assert params[2].default is None
    # Old two-argument path still renders the 12-sheet NIGHT-1 workbook.
    import tempfile
    from tools.build_excel_report_en import main
    assert main([BASE_RUN]) == 0
    assert _sha(NIGHT1_XLSX) == NIGHT1_SHA_AT_IMPORT
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "NIGHT1_REPORT_EN.xlsx")
        build_workbook(BASE_RUN, out)
        wb = _open_xlsx(out)
        assert wb.sheetnames == SHEET_NAMES


def test_scoped_filter_invariants():
    from bench import report as rep
    kept, excluded = _load_records_scoped(N2_RUN, NIGHT2_TAG_PREFIX)
    assert excluded == 720
    assert len(kept) >= 1 and kept
    assert all(str(rep.model_of(r)).startswith(NIGHT2_TAG_PREFIX)
               for r in kept)
    assert {str(rep.model_of(r)) for r in kept} == set(EXPECTED_N2_TAGS)
    raw_total = sum(1 for line in
                    open(os.path.join(N2_RUN, "results.jsonl"),
                         encoding="utf-8") if line.strip())
    assert len(kept) + excluded == raw_total
    assert len(kept) == 1090


def test_loader_invariants_and_stable_names():
    base_only = load_analysis(os.path.join(ROOT, "analysis_en"))
    assert base_only
    assert not [m for m in base_only
                if str(m.get("model", "")).startswith(NIGHT2_TAG_PREFIX)]
    base, n2 = load_analysis_28(os.path.join(ROOT, "analysis_en"))
    assert {str(m.get("model")) for m in n2} == set(EXPECTED_N2_TAGS)
    assert not ({str(m.get("model")) for m in base} & set(EXPECTED_N2_TAGS))
    assert GROUP_N2_FILENAME in GROUP_FILES_ALL
    for name in GROUP_FILES:
        assert name in GROUP_FILES_ALL
    assert os.path.isfile(os.path.join(ROOT, ELIGIBILITY_N2_RELPATH))
    import tools.build_excel_report_en as ben
    import tools.uk_sheets_deep as uk
    for name in ("normalize_evidence_quote", "strip_trailing_marker",
                 "ACCENT", "_style_table", "_q_scale"):
        assert hasattr(ben, name), name
    assert uk.normalize_quote_uk is not None
    # 28-model narrative builders render without new claims machinery.
    import openpyxl
    wb = openpyxl.Workbook()
    build_all_narrative_28(wb, base, n2, BASE_ID, N2_ID)
    assert wb.sheetnames[1:] == ["Model Scorecard", "Recommendations",
                                 "Strengths & Weaknesses"]


def test_no_exact_count_equalities_in_owned_code():
    rx = re.compile(r"(==|!=)\s*(24|4|576)\b")
    hits = []
    for path in (
            os.path.join(ROOT, "tools", "build_excel_report_en.py"),
            os.path.join(ROOT, "tools", "en_sheets_narrative.py"),
            os.path.join(ROOT, "tools", "en_sheets_detail.py"),
            os.path.join(ROOT, "tests", "test_report28_en28.py")):
        with open(path, encoding="utf-8") as f:
            for i, line in enumerate(f, 1):
                if rx.search(line):
                    hits.append("%s:%d:%s" % (path, i, line.strip()))
    assert hits == [], hits
