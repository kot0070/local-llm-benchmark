"""Tests for tools/build_excel_report.py (no network, no Ollama)."""
import csv
import hashlib
import json
import os
import sys

import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from tools.build_excel_report import (  # noqa: E402
    SHEET_NAMES,
    canonical_not_attempted_uk,
    main,
)

DIGEST = "abc123"

EXPECTED_SHEETS = [
    "Стислий підсумок",
    "Огляд",
    "Оцінка моделей",
    "Рекомендації",
    "Сильні та слабкі сторони",
    "Деталі по тестах",
    "Аналіз помилок",
    "HOME-вердикти",
    "Рейтинг по тестах",
    "Швидкість",
    "Відомі особливості",
    "HOME-07 (окремий прогін)",
]

RETIRED_SHEETS = ("Профілі моделей",)


def _rec(model, test, case, status, sem, sub=""):
    return {"key": "r|OLLAMA|%s@%s|%s@n1|%s|0|R0|p" % (model, DIGEST, test, case),
            "identity": {"tag": model, "digest": DIGEST,
                         "ollama_version": "0.34.2"},
            "response": {"content": "x", "thinking": "", "tool_calls": [],
                         "tool_call_channel": "none", "done_reason": "stop",
                         "prompt_eval_count": 1, "eval_count": 1},
            "timing": {"load_s": 0.1, "ttft_s": 0.2, "prompt_eval_s": 0.1,
                       "eval_s": 0.5, "prompt_tok_s": 10.0, "gen_tok_s": 2.0,
                       "wall_s": 1.0, "first_after_load": None},
            "resources": {"vram_peak_mb": 100.0, "vram_baseline_mb": 50.0,
                          "gpu_util_mean": 5.0, "temp_start": 40.0,
                          "temp_peak": 45.0, "throttle": False},
            "verdict": {"status": status, "sub_reason": sub,
                        "attribution": "MODEL", "sem": sem, "strict": sem,
                        "details": {}},
            "flags": [],
            "timestamps": {"utc": "2026-09-21T15:00:00+00:00",
                           "local": "2026-09-21T10:00:00-05:00"}}


def _write_run(run_dir, recs, with_perf=True, with_manifest=True):
    os.makedirs(run_dir, exist_ok=True)
    with open(os.path.join(run_dir, "results.jsonl"), "w",
              encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps(r) + "\n")
    if with_perf:
        with open(os.path.join(run_dir, "perf.csv"), "w", newline="",
                  encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["tag", "cold_load_s", "ttft_s", "prompt_tok_s",
                        "gen_tok_s", "vram_peak_mb", "offload_ratio",
                        "throttle"])
            w.writerow(["clean-ok:1b", 1.0, 0.2, 1000.0, 50.0, 500.0, 1.0,
                        False])
            w.writerow(["fail-model:7b", 9.0, 0.5, 800.0, 5.0, 7000, 0.5,
                        False])
    if with_manifest:
        man = {"run_id": os.path.basename(run_dir),
               "start_utc": "2026-09-21T14:00:00+00:00",
               "start_local": "2026-09-21T09:00:00-05:00",
               "budget_hours": 8,
               "ollama_version": "0.34.2",
               "fingerprint": {"gpu": "TEST-GPU", "cpu": "TEST-CPU",
                               "ram": "TEST-RAM", "os": "TEST-OS"}}
        with open(os.path.join(run_dir, "manifest.json"), "w",
                  encoding="utf-8") as f:
            json.dump(man, f)


def _synthetic_recs():
    recs = []
    for i in range(3):  # clean model, all OK
        recs.append(_rec("clean-ok:1b", "HOME-01", "HOME-01-c%d" % i,
                         "OK", 1.0))
    # failing model: 2 OK + 2 WRONG -> Q_sem 0.5 on HOME-03
    for i in range(2):
        recs.append(_rec("fail-model:7b", "HOME-03", "HOME-03-c%d" % i,
                         "OK", 1.0))
    for i in range(2, 4):
        recs.append(_rec("fail-model:7b", "HOME-03", "HOME-03-c%d" % i,
                         "WRONG_ANSWER", 0.0, sub="EVIDENCE_MISMATCH"))
    # UNSUPPORTED-only model (must show "поза рейтингом", no fabricated Q)
    for i in range(2):
        recs.append(_rec("unsup-model:1b", "HOME-02", "HOME-02-c%d" % i,
                         "UNSUPPORTED_CAPABILITY", 0.0, sub="LANG"))
    # NOT_RUN_BUDGET-only pair
    for i in range(2):
        recs.append(_rec("clean-ok:1b", "HOME-04", "HOME-04-c%d" % i,
                         "NOT_RUN_BUDGET", 0.0, sub="BUDGET"))
    return recs


def _open_xlsx(path):
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True)


def test_sheet_list_order_and_retired_absent(tmp_path):
    run_dir = str(tmp_path / "ORDER")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT.xlsx"))
    assert wb.sheetnames == EXPECTED_SHEETS
    assert SHEET_NAMES == EXPECTED_SHEETS
    assert len(SHEET_NAMES) == 12
    for retired in RETIRED_SHEETS:
        assert retired not in wb.sheetnames
    # the deep strengths sheet reuses the retired name exactly once
    assert wb.sheetnames.count("Сильні та слабкі сторони") == 1


def test_retired_builders_gone():
    import tools.build_excel_report as mod
    assert not hasattr(mod, "_sheet_profiles")
    assert not hasattr(mod, "_sheet_strengths")
    assert not hasattr(mod, "STRONG_HEADER")


def test_synthetic_sheets_values_and_determinism(tmp_path):
    from bench import report as rep
    run_dir = str(tmp_path / "SYNTH")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    out = os.path.join(run_dir, "NIGHT1_REPORT.xlsx")
    assert os.path.exists(out)
    with open(out, "rb") as f:
        first_bytes = f.read()
    assert main([run_dir]) == 0  # rerun: byte-identical output
    with open(out, "rb") as f:
        assert f.read() == first_bytes
    wb = _open_xlsx(out)
    assert wb.sheetnames == SHEET_NAMES
    # header row present and bold on every sheet
    for ws in wb.worksheets:
        assert ws.max_row >= 1 and ws.max_column >= 1
        assert ws["A1"].value not in (None, "")
        assert ws["A1"].font.bold is True
        assert ws.freeze_panes is not None
    # rating: fail-model HOME-03 Q_sem matches bench.report independently
    records = rep.load_records(run_dir)
    per_test, _, _ = rep.compute_tables(records)
    exp = next(r for r in per_test["HOME-03"]
               if r["model"] == "fail-model:7b")
    ws = wb["Рейтинг по тестах"]
    found = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == "HOME-03" and row[1] == "fail-model:7b":
            found = row
            break
    assert found is not None
    assert found[2] == pytest.approx(exp["q_sem"], abs=1e-9)
    # UNSUPPORTED-only row labelled, not ranked with a fabricated score
    ws_rows = list(ws.iter_rows(min_row=2, values_only=True))
    unsup = [r for r in ws_rows
             if r[0] == "HOME-02" and r[1] == "unsup-model:1b"]
    assert len(unsup) == 1
    assert unsup[0][-1] == "поза рейтингом"
    # NOT_RUN_BUDGET-only row labelled as well
    notrun = [r for r in ws_rows
              if r[0] == "HOME-04" and r[1] == "clean-ok:1b"]
    assert len(notrun) == 1
    assert notrun[0][-1] == "поза рейтингом"
    # verdicts sheet: one row per HOME test (24) + header
    wv = wb["HOME-вердикти"]
    assert wv.max_row == 25
    # HOME-07 sheet exists with the non-comparability note
    w7 = wb["HOME-07 (окремий прогін)"]
    texts = " ".join(str(c.value or "") for row in w7.iter_rows()
                     for c in row)
    assert "не порівнюються" in texts
    # deep sheets render from the shared analysis (24 models / 576 pairs)
    wd = wb["Деталі по тестах"]
    drows = list(wd.iter_rows(min_row=2, values_only=True))
    assert len(drows) == 576
    assert len({(r[0], r[1]) for r in drows}) == 576
    ws5 = wb["Сильні та слабкі сторони"]
    assert ws5.max_row == 25  # header + 24 model rows (deep, not template)


def test_summary_sheet_sections_and_numbers(tmp_path):
    run_dir = str(tmp_path / "SUMM")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT.xlsx"))
    assert wb.sheetnames[0] == "Стислий підсумок"
    ws = wb["Стислий підсумок"]
    assert ws["A1"].value == "Розділ"
    sections = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)
                if r[0]}
    assert list(sections) == ["Що це", "Масштаб прогону",
                              "Найкращі власні результати",
                              "Вердикти власних моделей", "Покриття часу",
                              "Форма проти змісту", "Відповідність залізу",
                              "Інженерна знахідка", "Глибина по моделях"]
    for title, text in sections.items():
        assert isinstance(text, str) and len(text) > 40, title
    # every section is Ukrainian prose, not a translation shell
    assert "аркушах" in sections["Глибина по моделях"]


def test_missing_gen_tok_s_writes_no_data(tmp_path):
    run_dir = str(tmp_path / "NODATA")
    recs = [_rec("embed-only:1b", "HOME-02", "HOME-02-c0",
                  "UNSUPPORTED_CAPABILITY", 0.0, sub="LANG")]
    for i in range(2):
        recs.append(_rec("text-model:1b", "HOME-01", "HOME-01-c%d" % i,
                          "OK", 1.0))
    os.makedirs(run_dir, exist_ok=True)
    with open(os.path.join(run_dir, "results.jsonl"), "w",
              encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps(r) + "\n")
    with open(os.path.join(run_dir, "perf.csv"), "w", newline="",
              encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["tag", "cold_load_s", "ttft_s", "prompt_tok_s",
                    "gen_tok_s", "vram_peak_mb", "offload_ratio",
                    "throttle"])
        w.writerow(["embed-only:1b", 1.0, 0.2, 1000.0, "", 500.0, 1.0,
                    False])
        w.writerow(["text-model:1b", 1.0, 0.2, 1000.0, 50.0, 500.0, 1.0,
                    False])
    man = {"run_id": "NODATA",
           "start_local": "2026-09-21T09:00:00-05:00",
           "budget_hours": 8}
    with open(os.path.join(run_dir, "manifest.json"), "w",
              encoding="utf-8") as f:
        json.dump(man, f)
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT.xlsx"))
    spd = wb["Швидкість"]
    spd_gen = next(c.column for c in spd[1]
                   if str(c.value or "") == "gen tok/s")
    emb_spd = next(r for r in spd.iter_rows(min_row=2, values_only=True)
                   if r[0] == "embed-only:1b")
    assert emb_spd[spd_gen - 1] == "немає даних"
    txt_spd = next(r for r in spd.iter_rows(min_row=2, values_only=True)
                   if r[0] == "text-model:1b")
    assert txt_spd[spd_gen - 1] == 50.0


def test_empty_results_no_crash(tmp_path):
    run_dir = str(tmp_path / "EMPTY")
    _write_run(run_dir, [], with_perf=False, with_manifest=False)
    assert main([run_dir]) == 0
    out = os.path.join(run_dir, "NIGHT1_REPORT.xlsx")
    assert os.path.exists(out)
    wb = _open_xlsx(out)
    assert wb.sheetnames == SHEET_NAMES


def test_canonical_not_attempted_uk_units():
    elig = {
        "HOME-05": {"m:1b": {"code": "U",
                             "reason": "CAP: немає vision"}},
        "HOME-01": {"m:1b": {"code": "S", "reason": "поза scope"}},
        "HOME-08": {"m:1b": {"code": "O", "reason": ""}},
        "HOME-18": {"m:1b": {"code": "O",
                             "reason": "багатоходовий діалог не "
                                       "задокументовано"}},
        "HOME-09": {"m:1b": {"code": "U",
                             "reason": "CONTEXT_WINDOW: потрібно 36000, "
                                       "runtime 8192"}},
    }
    # Identical causes render identically whatever the original phrasing.
    for note in ("Не виконувався: модель не має задокументованої візуальної "
                 "здатності, тож прогін зафіксував блок непідтримуваної "
                 "здатності.",
                 "Не виконувався: відсутня задокументована можливість, "
                 "немає підтримки бачення.",
                 "Не виконувався: модель недокументована для зорового "
                 "вводу."):
        assert canonical_not_attempted_uk("HOME-05", "m:1b", note, elig) == \
            "Не виконувався — у моделі немає задокументованої зорової " \
            "здатності для цього тесту."
    assert canonical_not_attempted_uk(
        "HOME-01", "m:1b",
        "Не виконувався: поза допустимим обсягом цього прогону для цієї "
        "моделі.", elig) == \
        "Не виконувався — поза запланованим обсягом цього прогону для " \
        "цієї моделі."
    assert canonical_not_attempted_uk(
        "HOME-08", "m:1b",
        "Не виконувався: цю пару не було заплановано в цьому прогоні.",
        elig) == "Не виконувався — пару не заплановано в цьому прогоні."
    assert canonical_not_attempted_uk(
        "HOME-18", "m:1b",
        "Не виконувався: багатоходові діалоги — недокументовані сценарії "
        "для цієї моделі.", elig) == \
        "Не виконувався — багатоходовий діалог не задокументований для " \
        "цієї моделі."
    assert canonical_not_attempted_uk(
        "HOME-09", "m:1b",
        "Не виконувався: тест потребує близько 36000 токенів контексту, а "
        "середовище дозволяє 8192.", elig) == \
        "Не виконувався — тесту потрібно 36000 токенів контексту, " \
        "середовище виконання має 8192."
    # Genuinely run-specific detail survives after the canonical clause.
    got = canonical_not_attempted_uk(
        "HOME-99", "m:1b",
        "Не виконувався: 8-годинний бюджет прогону вичерпався до цього "
        "блока.", {})
    assert got.startswith("Не виконувався — у цьому прогоні немає оцінених "
                          "кейсів. ")
    assert "бюджет" in got
    # Retired prefixes never appear.
    for got2 in (canonical_not_attempted_uk("HOME-05", "m:1b", "x", elig),
                 got):
        assert "Не оцінювався" not in got2
        assert got2.startswith("Не виконувався — ")


def test_quote_normalization_reuses_shared_helper():
    from tools.build_excel_report import normalize_failure_quotes_uk
    from tools.build_excel_report_en import (
        TRUNCATION_MARKER,
        UNVERIFIABLE_QUOTE,
        normalize_evidence_quote,
    )
    recs = [_rec("m:1b", "HOME-01", "HOME-01-c0", "WRONG_ANSWER", 0.0)]
    recs[0]["response"]["content"] = "повна відповідь моделі тут"
    analysis = [{"model": "m:1b",
                 "failure_analysis": [
                     {"pattern": "WRONG_ANSWER[x]",
                      "example_case_id": "HOME-01-c0",
                      "example_quote": "повна відповідь моделі тут"},
                     {"pattern": "WRONG_ANSWER[y]",
                      "example_case_id": "HOME-01-c0",
                      "example_quote": "повна відповідь"},
                     {"pattern": "WRONG_ANSWER[z]",
                      "example_case_id": "HOME-01-c0",
                      "example_quote": "чужий текст"},
                 ]}]
    counts = normalize_failure_quotes_uk(analysis, recs)
    assert counts == {"complete": 1, "truncated": 1, "unverifiable": 1}
    fa = analysis[0]["failure_analysis"]
    assert fa[0]["example_quote"] == "повна відповідь моделі тут"
    assert fa[1]["example_quote"] == "повна відповідь" + TRUNCATION_MARKER
    assert fa[2]["example_quote"] == UNVERIFIABLE_QUOTE
    # same derivation as the English helper, case for case
    assert normalize_evidence_quote("повна відповідь",
                                    "повна відповідь моделі тут") == \
        ("повна відповідь" + TRUNCATION_MARKER, "truncated")


REAL_RUN = os.path.join(ROOT, "results", "night_20260921-155146")


def test_detail_sheet_canonical_notes_real_run():
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    assert "Оцінка моделей" in wb.sheetnames
    assert "Рекомендації" in wb.sheetnames
    assert "Сильні та слабкі сторони" in wb.sheetnames
    ws = wb["Деталі по тестах"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 576
    assert len({(r[0], r[1]) for r in rows}) == 576
    assert len({r[0] for r in rows}) == 24
    nas = [r for r in rows if r[5] == "не виконувався"]
    assert len(nas) == 361
    for r in nas:
        assert r[3] is None, r[:2]  # Q_sem genuinely empty, never 0
        assert isinstance(r[6], str) and \
            r[6].startswith("Не виконувався — "), r[:2]
    for r in rows:
        note = r[6] if isinstance(r[6], str) else ""
        assert "Не оцінювався" not in note, r[:2]
        assert "Ніколи" not in note, r[:2]
    # vocabulary is short: canonical templates, not 40 drifting phrasings
    distinct = {r[6] for r in nas}
    assert len(distinct) <= 20, len(distinct)


def test_failure_quotes_verified_real_run():
    from bench import report as rep
    from tools.build_excel_report_en import (
        TRUNCATION_MARKER,
        UNVERIFIABLE_QUOTE,
        strip_trailing_marker,
    )
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ws = wb["Аналіз помилок"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 133
    assert len({r[0] for r in rows}) == 24
    records = rep.load_records(REAL_RUN)
    index = {}
    for r in records:
        tag = rep.model_of(r)
        case = str(r.get("key", "")).split("|")[4]
        index.setdefault((tag, case),
                         (r.get("response") or {}).get("content", ""))
    complete = truncated = quoted = 0
    for r in rows:
        q = r[6]
        if q in (None, "Немає даних") or \
                (isinstance(q, str) and q.startswith("Жоден шаблон")):
            continue
        quoted += 1
        assert q != UNVERIFIABLE_QUOTE, (r[0], r[5])
        assert "[...]" not in q, (r[0], r[5])
        stripped = strip_trailing_marker(q)
        assert "[truncated]" not in stripped, (r[0], r[5])
        full = index.get((r[0], r[5]))
        assert full is not None, (r[0], r[5])
        if stripped == full:
            complete += 1
        else:
            assert full.startswith(stripped) and stripped, (r[0], r[5])
            assert q.endswith(TRUNCATION_MARKER), (r[0], r[5])
            truncated += 1
    assert quoted == 110
    assert complete == 53
    assert truncated == 57


def test_cross_workbook_numbers_match_english():
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT.xlsx")
    en_out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not (os.path.exists(out) and os.path.exists(en_out)):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    en = _open_xlsx(en_out)
    # ratings equal per model
    enm = {r[0]: r[1] for r in en["Model Scorecard"].iter_rows(
        values_only=True)
        if isinstance(r[0], str) and isinstance(r[1], (int, float))}
    ukm = {r[0]: r[1] for r in wb["Оцінка моделей"].iter_rows(
        values_only=True)
        if isinstance(r[0], str) and isinstance(r[1], (int, float))}
    assert set(enm) == set(ukm) and len(enm) == 24
    for tag, v in enm.items():
        assert ukm[tag] == v, tag
    # Q_sem and case counts equal per (model, test): not-attempted rows
    # carry empty cells in both workbooks (FINAL_POLISH fix 1), and the
    # former gemma3:12b/HOME-02 JSON mismatch is corrected at the source
    # (analysis_uk/group_a.json n 0/30 -> 0/40, FINAL_POLISH fix 2).
    end = {(r[0], r[1]): (r[3], r[4]) for r in
           en["Per-Model Test Detail"].iter_rows(min_row=2,
                                                 values_only=True)}
    ukd = {(r[0], r[1]): (r[3], r[4]) for r in
           wb["Деталі по тестах"].iter_rows(min_row=2, values_only=True)}
    assert set(end) == set(ukd)
    for k, (q, n) in end.items():
        uq, un = ukd[k]
        assert (uq, un) == (q, n), k


def test_spotcheck_audited_numbers_unchanged():
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ov = {r[0]: r[1] for r in wb["Огляд"].iter_rows(values_only=True)
          if r[0]}
    assert ov["run_id"] == "night_20260921-155146"
    assert ov["моделей запущено"] == 24
    assert ov["тестів запущено"] == 25
    assert ov["записів"] == 4190
    assert ov["NOT_RUN_BUDGET"] == 799
    assert ov["NOT_RUN_BUDGET, %"] == 19.1
    hv = {r[0]: r for r in wb["HOME-вердикти"].iter_rows(min_row=2,
                                                        values_only=True)}
    assert hv["HOME-01"][3] == 0.858  # home Q_sem
    assert hv["HOME-01"][6] == 0.944  # best competitor Q_sem
    assert hv["HOME-03"][3] == 0  # command-r7b home zero
    fr = list(wb["Рейтинг по тестах"].iter_rows(min_row=2, values_only=True))
    assert len(fr) == 576
    top = next(r for r in fr if r[0] == "HOME-01" and r[1] == "qwen3:8b")
    assert (top[2], top[3]) == (1, 1)
    pf = {r[0]: r for r in wb["Швидкість"].iter_rows(min_row=2,
                                                    values_only=True)
          if r[0]}
    assert pf["functiongemma:270m"][4] == pytest.approx(460.974, abs=0.01)
    h7 = list(wb["HOME-07 (окремий прогін)"].iter_rows(values_only=True))
    assert (h7[1][1], h7[1][2]) == (0, 0)
    assert (h7[3][1], h7[3][2]) == ("16/16", "16/17")
    ex = {r[0]: r[1] for r in wb["Стислий підсумок"].iter_rows(
        min_row=2, values_only=True)}
    assert "0 перемог" in ex["Вердикти власних моделей"]
    assert "12 нічиїх" in ex["Вердикти власних моделей"]
    assert "9 поразок" in ex["Вердикти власних моделей"]


def test_determinism_sha(tmp_path):
    run_dir = str(tmp_path / "DET")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    out = os.path.join(run_dir, "NIGHT1_REPORT.xlsx")
    with open(out, "rb") as f:
        first = f.read()
    assert main([run_dir]) == 0
    with open(out, "rb") as f:
        second = f.read()
    assert first == second
    assert hashlib.sha256(first).hexdigest() == \
        hashlib.sha256(second).hexdigest()


def test_real_run_exit_0_and_readonly():
    if not os.path.isdir(REAL_RUN):
        pytest.skip("missing real run dir")
    before = set(os.listdir(REAL_RUN))
    assert main(["night_20260921-155146"]) == 0
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT.xlsx")
    assert os.path.exists(out)
    wb = _open_xlsx(out)
    assert wb.sheetnames == SHEET_NAMES
    # tool writes only its own declared output file
    new_files = set(os.listdir(REAL_RUN)) - before
    assert new_files <= {"NIGHT1_REPORT.xlsx"}, new_files


def test_no_gpu_uuid_in_any_workbook_cell(tmp_path):
    import re
    uuid_pat = re.compile(r"GPU-[0-9a-f]{8}-")
    # synthetic workbook: no fingerprint-derived dump may leak through
    run_dir = str(tmp_path / "NOGPUUID")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT.xlsx"))
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    assert not uuid_pat.search(v), (ws.title, v)
    # real workbooks (both languages), when present
    for name in ("NIGHT1_REPORT.xlsx", "NIGHT1_REPORT_EN.xlsx"):
        path = os.path.join(REAL_RUN, name)
        if not os.path.exists(path):
            continue
        wb = _open_xlsx(path)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        assert not uuid_pat.search(v), (name, ws.title, v)


def test_overview_gpu_cell_rendered_without_uuid(tmp_path):
    run_dir = str(tmp_path / "GPUCELL")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT.xlsx"))
    ov = {r[0]: r[1] for r in wb["Огляд"].iter_rows(values_only=True)
          if r[0]}
    assert ov["GPU"] == "NVIDIA GeForce RTX 3070 (8 ГБ VRAM)"
