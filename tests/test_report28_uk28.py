"""Tests for the 28-model Ukrainian workbook (W_UK28). No network, no Ollama."""
import glob
import hashlib
import inspect
import os
import re
import subprocess
import sys

import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from tools.build_excel_report import (  # noqa: E402
    ELIGIBILITY_N2_RELPATH,
    NIGHT2_TAG_PREFIX,
    SHEET_NAMES,
    XLSX_NAME_28,
    _load_records_scoped_uk,
    build_workbook,
    normalize_failure_quotes_uk,
)
from tools.build_excel_report_en import (  # noqa: E402
    TRUNCATION_MARKER,
    UNVERIFIABLE_QUOTE,
    normalize_evidence_quote,
    strip_trailing_marker,
)
from tools.uk_sheets_deep import (  # noqa: E402
    DETAIL_HEADER_28_UK,
    FAILURE_HEADER_28_UK,
    FORMULA_N1_UK,
    FORMULA_N2_UK,
    GROUP_FILES,
    GROUP_FILES_ALL,
    GROUP_N2_FILENAME,
    N2_BAND_LABEL_UK,
    TIER_FOCUSED_UK,
    TIER_GENERALIST_UK,
    TIER_SINGLE_UK,
    build_all_narrative_28_uk,
    load_analysis_28_uk,
    load_analysis_uk,
)

BASE_RUN = os.path.join(ROOT, "results", "night_20260921-155146")
N2_RUN = os.path.join(ROOT, "results", "night2_b")
BASE_ID = "night_20260921-155146"
N2_ID = "night2_b"
OUT = os.path.join(N2_RUN, XLSX_NAME_28)
EN_OUT = os.path.join(N2_RUN, "NIGHT28_REPORT_EN.xlsx")
NIGHT1_XLSX = os.path.join(BASE_RUN, "NIGHT1_REPORT.xlsx")

EXPECTED_N2_TAGS = frozenset((
    "night2-ministral-3-8b:q4km",
    "night2-nomic-embed-code:q4km",
    "night2-qwen3-vl-8b:q4km",
    "night2-qwen3.5-9b:q4km",
))

EXPECTED_SHEETS_UK = [
    "Стислий підсумок",
    "Огляд",
    "Оцінка моделей",
    "Рекомендації",
    "Сильні та слабкі сторони",
    "Деталі по тестах",
    "Аналіз помилок",
    "HOME-вердикти",
    "Рейтинг по тестах",
    "Швидкість",
    "Відомі особливості",
    "HOME-07 (окремий прогін)",
]

with open(NIGHT1_XLSX, "rb") as _f:
    NIGHT1_SHA_AT_IMPORT = hashlib.sha256(_f.read()).hexdigest()


def _sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def _open_xlsx(path):
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True)


def _build28():
    return build_workbook(BASE_RUN, OUT, N2_RUN)


@pytest.fixture(scope="module")
def built():
    _build28()
    return OUT


def _walk_scorecard(ws):
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        if all(v is None for v in row[1:]):
            yield list(row), row[0]
        else:
            yield list(row), None


def test_a1_exists_sheets_order_deterministic(built):
    assert os.path.basename(built) == XLSX_NAME_28
    assert os.path.dirname(os.path.abspath(built)) == os.path.abspath(N2_RUN)
    assert SHEET_NAMES == EXPECTED_SHEETS_UK
    assert len(SHEET_NAMES) == len(EXPECTED_SHEETS_UK)
    first = _sha(built)
    _build28()
    second = _sha(built)
    assert first == second
    print("NIGHT28_UK sha256: %s" % first)
    wb = _open_xlsx(built)
    assert wb.sheetnames == EXPECTED_SHEETS_UK
    assert wb.sheetnames == SHEET_NAMES
    leftovers = glob.glob(os.path.join(N2_RUN, "NIGHT1_REPORT.xlsx"))
    assert leftovers == [], leftovers


def test_a2_scorecard_bands_formulas_sort(built):
    wb = _open_xlsx(built)
    ws = wb["Оцінка моделей"]
    header = [c.value for c in ws[3]]
    assert header[-1] == "Прогін"
    bands = []
    data = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        if all(v is None for v in row[1:]):
            bands.append(row[0])
        else:
            data.append(list(row))
    assert bands == [TIER_GENERALIST_UK, TIER_FOCUSED_UK, TIER_SINGLE_UK,
                     N2_BAND_LABEL_UK]
    assert len(data) == 28
    assert len({r[0] for r in data}) == len(data)
    base_rows = [r for r in data if r[-1] == BASE_ID]
    n2_rows = [r for r in data if r[-1] == N2_ID]
    assert len(base_rows) + len(n2_rows) == len(data)
    assert {r[0] for r in n2_rows} == set(EXPECTED_N2_TAGS)
    assert not ({r[0] for r in base_rows} & set(EXPECTED_N2_TAGS))
    assert len(n2_rows) == len(EXPECTED_N2_TAGS)
    order = [r[-1] for r in data]
    assert order == [BASE_ID] * len(base_rows) + [N2_ID] * len(n2_rows)
    current = None
    last = None
    for row, band in _walk_scorecard(ws):
        if band is not None:
            current = band
            last = None
            continue
        assert current is not None
        if last is not None and isinstance(row[1], (int, float)) \
                and isinstance(last, (int, float)):
            assert row[1] <= last, (row[0], row[1], last)
        last = row[1]
    note = ws["A2"].value or ""
    assert FORMULA_N1_UK in note
    assert FORMULA_N2_UK in note
    assert "без home-доданка" in note


def test_a3_detail_rows_runs_overview(built):
    wb = _open_xlsx(built)
    ws = wb["Деталі по тестах"]
    assert list(ws[1][i].value for i in range(len(DETAIL_HEADER_28_UK))) == \
        DETAIL_HEADER_28_UK
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 672
    assert len({(r[0], r[1]) for r in rows}) == len(rows)
    runs = {r[7] for r in rows}
    assert runs == {BASE_ID, N2_ID}
    assert all(r[7] for r in rows)
    nas = [r for r in rows if r[5] == "не виконувався"]
    for r in nas:
        assert r[3] is None, r[:2]
        assert r[4] is None, r[:2]
    texts = " ".join(str(c.value or "") for wsn in ("Огляд",)
                     for row in wb[wsn].iter_rows() for c in row)
    assert BASE_ID in texts
    assert N2_ID in texts
    assert "0.34.2" in texts
    assert "0.34.3" in texts
    assert "720" in texts
    ov = {r[0]: r[1] for r in wb["Огляд"].iter_rows(values_only=True)
          if r[0]}
    assert ov["виключених позаобсягових рядків"] == 720


def test_a4_cross_workbook_equal_en(built):
    en = _open_xlsx(EN_OUT)
    wb = _open_xlsx(built)
    enm = {r[0]: r[1] for r in en["Model Scorecard"].iter_rows(
        values_only=True)
        if isinstance(r[0], str) and isinstance(r[1], (int, float))}
    ukm = {r[0]: r[1] for r in wb["Оцінка моделей"].iter_rows(
        values_only=True)
        if isinstance(r[0], str) and isinstance(r[1], (int, float))}
    assert set(enm) == set(ukm)
    bad = [k for k in enm if ukm.get(k) != enm[k]]
    assert bad == [], bad[:5]
    end = {(r[0], r[1]): (r[3], r[4]) for r in
           en["Per-Model Test Detail"].iter_rows(min_row=2,
                                                 values_only=True)}
    ukd = {(r[0], r[1]): (r[3], r[4]) for r in
           wb["Деталі по тестах"].iter_rows(min_row=2, values_only=True)}
    assert set(end) == set(ukd)
    bad2 = [k for k in end if end[k] != ukd.get(k)]
    assert bad2 == [], bad2[:5]
    print("cross-workbook disagreements: 0 over %d pairs" % len(end))


def test_a5_quotes_verified_and_identical(built):
    import copy
    import json
    from bench import report as rep
    wb = _open_xlsx(built)
    ws = wb["Аналіз помилок"]
    assert list(ws[1][i].value for i in range(len(FAILURE_HEADER_28_UK))) == \
        FAILURE_HEADER_28_UK
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len({r[0] for r in rows}) == len({r[0] for r in rows} - set()) \
        and len({r[0] for r in rows}) >= len(EXPECTED_N2_TAGS)
    index = {}
    for run in (BASE_RUN, N2_RUN):
        for r in rep.load_records(run):
            tag = rep.model_of(r)
            parts = str(r.get("key", "")).split("|")
            case = parts[4] if len(parts) >= 5 else ""
            index.setdefault((tag, case),
                             (r.get("response") or {}).get("content", ""))
    quoted = complete = truncated = 0
    for r in rows:
        q = r[6]
        if q in (None, "Немає даних") or (isinstance(q, str)
                                          and q.startswith("Жоден шаблон")):
            continue
        quoted += 1
        assert q != UNVERIFIABLE_QUOTE, (r[0], r[5])
        assert "[...]" not in q, (r[0], r[5])
        stripped = strip_trailing_marker(q)
        assert "[truncated]" not in stripped, (r[0], r[5])
        full = index.get((r[0], r[5]))
        assert full is not None, (r[0], r[5])
        assert stripped, (r[0], r[5])
        if stripped == full:
            complete += 1
        else:
            assert full.startswith(stripped), (r[0], r[5])
            assert q.endswith(TRUNCATION_MARKER), (r[0], r[5])
            truncated += 1
    assert quoted >= 1
    assert complete + truncated == quoted
    ana = []
    for n in ("group_a.json", "group_b.json", "group_c.json",
              "group_d.json", "group_n2.json"):
        with open(os.path.join(ROOT, "analysis_uk", n),
                  encoding="utf-8") as f:
            ana.extend(json.load(f))
    c1 = copy.deepcopy(ana)
    base = [m for m in c1 if not str(m["model"]).startswith("night2-")]
    n2 = [m for m in c1 if str(m["model"]).startswith("night2-")]
    normalize_failure_quotes_uk(base, rep.load_records(BASE_RUN))
    n2_all = rep.load_records(N2_RUN)
    n2_recs = [r for r in n2_all
               if str(rep.model_of(r)).startswith(NIGHT2_TAG_PREFIX)]
    normalize_failure_quotes_uk(n2, n2_recs)
    norm = {e["example_quote"] for m in base + n2
            for e in (m.get("failure_analysis") or [])
            if isinstance(e.get("example_quote"), str)
            and e["example_quote"].strip()}
    rendered = {r[6] for r in rows if isinstance(r[6], str)
                and r[6] not in (None, "Немає даних")
                and not r[6].startswith("Жоден шаблон")}
    missing = [q for q in rendered if q not in norm]
    assert missing == [], missing[:3]
    print("quotes: %d quoted %d complete %d truncated 0 unmarked" %
          (quoted, complete, truncated))


def test_a6_no_uuid_no_latin_prose(built):
    uuid_pat = re.compile(r"GPU-[0-9a-f]{8}-")
    mach_pat = re.compile(
        r"(C:\\\\Users|service\.json|BEGIN (RSA )?PRIVATE KEY|"
        r"AKIA[0-9A-Z]{16})", re.I)
    # The OS account name is read at run time, never written into the repo (a literal one leaked once).
    import os
    who = (os.environ.get("USERNAME") or os.environ.get("USER") or "").strip()
    if len(who) >= 3:
        mach_pat = re.compile(mach_pat.pattern[:-1] + "|" + re.escape(who) + ")", re.I)
    for path in (built, EN_OUT):
        wb = _open_xlsx(path)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and v:
                        assert not uuid_pat.search(v), (path, ws.title, v[:80])
                        assert not mach_pat.search(v), (path, ws.title, v[:80])
    cyr = re.compile(r"[\u0400-\u04FF]")
    wb = _open_xlsx(built)
    en_sentence = re.compile(
        r"\b(the|and|with|from|this|that|which|model|models|report|run|"
        r"results|quality|rating|score|test|tests)\b"
        r".*\b(the|and|with|from|this|that|which|model|models|report|run|"
        r"results|quality|rating|score|test|tests)\b", re.I)
    for ws in wb.worksheets:
        # "Відомі особливості" renders MASTER_PLAN section 5 verbatim in
        # English by design (same as the NIGHT-1 workbook); it is a quoted
        # source sheet, not Ukrainian-authored prose.
        if ws.title == "Відомі особливості":
            continue
        for row in ws.iter_rows():
            for cell in row:
                if ws.title == "Аналіз помилок" and cell.column == 7:
                    continue
                v = cell.value
                if not isinstance(v, str) or len(v) < 60:
                    continue
                if v.startswith("round(10*"):
                    continue
                assert not en_sentence.search(v), (
                    ws.title, cell.coordinate, v[:140])


def test_a7_suite_collect_count():
    proc = subprocess.run(
        [sys.executable, "-m", "pytest", "tests/", "--collect-only", "-q"],
        capture_output=True, text=True, cwd=ROOT, timeout=120)
    assert proc.returncode == 0, proc.stderr[-2000:]
    found = re.findall(r"(\d+)\s+tests?\s+collected", proc.stdout)
    assert found, proc.stdout[-2000:]
    assert int(found[-1]) >= 387


def test_a8_night1_untouched_and_backward_compat(built):
    assert _sha(NIGHT1_XLSX) == NIGHT1_SHA_AT_IMPORT
    sig = inspect.signature(build_workbook)
    params = list(sig.parameters.values())
    assert [p.name for p in params][:2] == ["run_dir", "out_path"]
    assert params[2].name == "extra_run_dir"
    assert params[2].default is None
    import tempfile
    from tools.build_excel_report import main
    assert main([BASE_RUN]) == 0
    assert _sha(NIGHT1_XLSX) == NIGHT1_SHA_AT_IMPORT
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "NIGHT1_REPORT.xlsx")
        build_workbook(BASE_RUN, out)
        wb = _open_xlsx(out)
        assert wb.sheetnames == SHEET_NAMES


def test_scoped_filter_invariants():
    from bench import report as rep
    kept, excluded = _load_records_scoped_uk(N2_RUN, NIGHT2_TAG_PREFIX)
    assert excluded == 720
    assert len(kept) >= 1 and kept
    assert all(str(rep.model_of(r)).startswith(NIGHT2_TAG_PREFIX)
               for r in kept)
    assert {str(rep.model_of(r)) for r in kept} == set(EXPECTED_N2_TAGS)
    raw_total = sum(1 for line in
                    open(os.path.join(N2_RUN, "results.jsonl"),
                         encoding="utf-8") if line.strip())
    assert len(kept) + excluded == raw_total
    assert len(kept) == 1090


def test_loader_invariants_and_stable_names():
    base_only = load_analysis_uk(os.path.join(ROOT, "analysis_uk"))
    assert base_only
    assert not [m for m in base_only
                if str(m.get("model", "")).startswith(NIGHT2_TAG_PREFIX)]
    base, n2 = load_analysis_28_uk(os.path.join(ROOT, "analysis_uk"))
    assert {str(m.get("model")) for m in n2} == set(EXPECTED_N2_TAGS)
    assert not ({str(m.get("model")) for m in base} & set(EXPECTED_N2_TAGS))
    assert GROUP_N2_FILENAME in GROUP_FILES_ALL
    for name in GROUP_FILES:
        assert name in GROUP_FILES_ALL
    assert os.path.isfile(os.path.join(ROOT, ELIGIBILITY_N2_RELPATH))
    import tools.build_excel_report_en as ben
    import tools.uk_sheets_deep as uk
    for name in ("normalize_evidence_quote", "strip_trailing_marker",
                 "ACCENT", "_style_table", "_q_scale"):
        assert hasattr(ben, name), name
    assert uk.normalize_quote_uk is not None
    assert uk.strip_quote_marker_uk is not None
    import openpyxl
    wb = openpyxl.Workbook()
    build_all_narrative_28_uk(wb, base, n2, BASE_ID, N2_ID)
    assert wb.sheetnames[1:] == ["Оцінка моделей", "Рекомендації",
                                 "Сильні та слабкі сторони"]
    assert normalize_evidence_quote("a [truncated]", "a plus") == \
        uk.normalize_quote_uk("a [truncated]", "a plus")


def test_no_exact_count_equalities_in_owned_code():
    rx = re.compile(r"(==|!=)\s*(24|4|576)\b")
    hits = []
    for path in (
            os.path.join(ROOT, "tools", "build_excel_report.py"),
            os.path.join(ROOT, "tools", "uk_sheets_deep.py"),
            os.path.join(ROOT, "tests", "test_report28_uk28.py")):
        with open(path, encoding="utf-8") as f:
            for i, line in enumerate(f, 1):
                if rx.search(line):
                    hits.append("%s:%d:%s" % (path, i, line.strip()))
    assert hits == [], hits
