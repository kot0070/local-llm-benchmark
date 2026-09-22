"""Tests for tools/uk_sheets_deep.py (Ukrainian deep sheets)."""
from __future__ import annotations

import json
import os

import openpyxl
import pytest

from tools import uk_sheets_deep as uk

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
ANALYSIS_DIR = os.path.join(ROOT, "analysis_uk")


def _load_groups() -> list[dict]:
    models: list[dict] = []
    for name in uk.GROUP_FILES:
        with open(os.path.join(ANALYSIS_DIR, name), encoding="utf-8") as f:
            models.extend(json.load(f))
    return models


def _build() -> tuple[object, list[dict]]:
    analysis = uk.load_analysis_uk(ANALYSIS_DIR)
    wb = openpyxl.Workbook()
    uk.build_all_uk(wb, analysis)
    return wb, analysis


def _sheet_values(wb, title):
    ws = wb[title]
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_loader_merges_24_sorted_models():
    analysis = uk.load_analysis_uk(ANALYSIS_DIR)
    assert len(analysis) == 24
    tags = [m["model"] for m in analysis]
    assert tags == sorted(tags)
    assert len(set(tags)) == 24


def test_loader_rejects_missing_dir():
    with pytest.raises(ValueError):
        uk.load_analysis_uk(os.path.join(ROOT, "analysis_uk_missing_xyz"))


def test_sheet_names_and_order():
    wb, _ = _build()
    assert wb.sheetnames[1:] == uk.SHEET_NAMES_UK
    assert uk.SHEET_NAMES_UK == [
        "Оцінка моделей",
        "Рекомендації",
        "Сильні та слабкі сторони",
        "Деталі по тестах",
        "Аналіз помилок",
    ]


def test_per_model_sheets_have_24_rows_each():
    wb, analysis = _build()
    tags = {str(m["model"]) for m in analysis}
    score_rows = _sheet_values(wb, "Оцінка моделей")
    model_rows = [r for r in score_rows if r[0] in tags]
    assert len(model_rows) == 24
    assert {r[0] for r in model_rows} == tags
    for title in ("Рекомендації", "Сильні та слабкі сторони"):
        rows = _sheet_values(wb, title)
        assert len(rows) - 1 == 24, title
        assert {r[0] for r in rows[1:]} == tags, title


def test_detail_576_unique_pairs():
    wb, analysis = _build()
    rows = _sheet_values(wb, "Деталі по тестах")
    data = rows[1:]
    assert len(data) == 576
    pairs = [(r[0], r[1]) for r in data]
    assert len(set(pairs)) == 576
    tags = {str(m["model"]) for m in analysis}
    assert {p[0] for p in pairs} == tags
    from collections import Counter
    per_model = Counter(p[0] for p in pairs)
    assert set(per_model.values()) == {24}


def test_failure_sheet_covers_all_24_models():
    wb, analysis = _build()
    rows = _sheet_values(wb, "Аналіз помилок")
    tags = {str(m["model"]) for m in analysis}
    present = {r[0] for r in rows[1:]}
    assert tags <= present
    from collections import Counter
    per_model = Counter(r[0] for r in rows[1:])
    assert all(v >= 1 for v in per_model.values())
    # total failure rows equals the JSON pattern count (all models have >=1)
    expected = sum(len(m.get("failure_analysis") or []) for m in analysis)
    assert len(rows) - 1 == expected


def test_rating_range_and_tier_order():
    wb, analysis = _build()
    rows = _sheet_values(wb, "Оцінка моделей")
    tags = {str(m["model"]) for m in analysis}
    bands = [uk.TIER_GENERALIST_UK, uk.TIER_FOCUSED_UK, uk.TIER_SINGLE_UK]
    seen_bands: list[str] = []
    current: str | None = None
    ratings: dict[str, list[float]] = {}
    for r in rows:
        if r[0] in bands:
            current = r[0]
            seen_bands.append(current)
        elif r[0] in tags:
            assert current is not None
            v = r[1]
            assert isinstance(v, (int, float)) and 0.0 <= v <= 10.0
            ratings.setdefault(current, []).append(float(v))
    assert seen_bands == bands
    for band, vals in ratings.items():
        assert vals == sorted(vals, reverse=True), band
    # tier membership matches tests_scored thresholds
    by_tag = {str(m["model"]): m for m in analysis}
    band_of_row: dict[str, str] = {}
    current2: str | None = None
    for r in rows:
        if r[0] in bands:
            current2 = r[0]
        elif r[0] in tags:
            band_of_row[r[0]] = current2
    for tag, band in band_of_row.items():
        expected_band, _ = uk._tier_of(by_tag[tag]["overall"]["tests_scored"])
        assert band == expected_band, tag
    # explicit: single-test band holds the 1-test specialists only
    singles = [r[0] for r in rows if r[0] in tags
               and by_tag[r[0]]["overall"]["tests_scored"] == 1]
    assert len(singles) == sum(
        1 for m in analysis if m["overall"]["tests_scored"] == 1)


def test_not_attempted_empty_quality_and_reason():
    wb, _ = _build()
    rows = _sheet_values(wb, "Деталі по тестах")
    na = [r for r in rows[1:] if r[5] == "не виконувався"]
    assert len(na) == 361
    for r in na:
        assert r[3] is None  # genuinely empty, not a misleading 0
        assert isinstance(r[6], str) and r[6].strip()
    scored = [r for r in rows[1:] if r[5] != "не виконувався"]
    assert len(scored) == 215
    for r in scored:
        assert isinstance(r[3], (int, float))


def test_not_attempted_empty_cases_and_quality():
    wb, _ = _build()
    rows = _sheet_values(wb, "Деталі по тестах")
    na = [r for r in rows[1:] if r[5] == "не виконувався"]
    assert len(na) == 361, len(na)
    for r in na:
        assert r[3] is None, r[:2]  # Q_sem genuinely empty, not a misleading 0
        assert r[4] is None, r[:2]  # Cases empty: no attempt count exists
    scored = [r for r in rows[1:] if r[5] != "не виконувався"]
    assert len(scored) == 215, len(scored)
    for r in scored:
        assert r[3] is not None, r[:2]
        assert r[4] is not None and str(r[4]).strip(), r[:2]


def test_quotes_byte_identical_to_analysis():
    wb, _ = _build()
    rows = _sheet_values(wb, "Аналіз помилок")
    rendered = [r[6] for r in rows[1:]]
    for m in _load_groups():
        for e in m.get("failure_analysis") or []:
            q = e.get("example_quote")
            if isinstance(q, str) and q.strip():
                assert q in rendered, (m["model"], e.get("example_case_id"))


def test_outcome_vocabulary_only_five():
    wb, _ = _build()
    rows = _sheet_values(wb, "Деталі по тестах")
    allowed = {"сильно", "задовільно", "слабко", "провал", "не виконувався"}
    assert uk.OUTCOME_UK == tuple(sorted(allowed, key=list(allowed).index)) \
        or set(uk.OUTCOME_UK) == allowed
    for r in rows[1:]:
        assert r[5] in allowed, r[5]


@pytest.mark.parametrize("n,expected", [
    (1, "1 раз"), (2, "2 рази"), (5, "5 разів"), (11, "11 разів"),
    (12, "12 разів"), (14, "14 разів"), (21, "21 раз"),
    (22, "22 рази"), (25, "25 разів"), (91, "91 раз"),
])
def test_uk_raziv_boundaries(n, expected):
    assert uk.uk_raziv(n) == expected


@pytest.mark.parametrize("n,expected", [
    (1, "1 запис"), (2, "2 записи"), (5, "5 записів"), (11, "11 записів"),
    (12, "12 записів"), (14, "14 записів"), (21, "21 запис"),
    (22, "22 записи"), (25, "25 записів"), (91, "91 запис"),
])
def test_uk_zapisiv_boundaries(n, expected):
    assert uk.uk_zapisiv(n) == expected


@pytest.mark.parametrize("n,expected", [
    (1, "1 випадок"), (2, "2 випадки"), (5, "5 випадків"),
    (11, "11 випадків"), (12, "12 випадків"), (14, "14 випадків"),
    (21, "21 випадок"), (22, "22 випадки"), (25, "25 випадків"),
    (91, "91 випадок"),
])
def test_uk_vipadkiv_boundaries(n, expected):
    assert uk.uk_vipadkiv(n) == expected


def test_quote_helpers_reuse_shared_implementation():
    from tools.build_excel_report_en import (
        TRUNCATION_MARKER,
        UNVERIFIABLE_QUOTE,
        normalize_evidence_quote,
        strip_trailing_marker,
    )
    # strip delegation handles every observed marker convention
    assert uk.strip_quote_marker_uk("текст [truncated]") == "текст"
    assert uk.strip_quote_marker_uk("текст [...]") == "текст"
    assert uk.strip_quote_marker_uk("текст без маркера") == \
        "текст без маркера"
    assert strip_trailing_marker("x [truncated]") == \
        uk.strip_quote_marker_uk("x [truncated]")
    # normalize delegation: proper prefix -> single consistent marker
    rendered, status = uk.normalize_quote_uk("префікс [truncated]",
                                             "префікс і решта")
    assert (rendered, status) == ("префікс" + TRUNCATION_MARKER, "truncated")
    assert normalize_evidence_quote("a [truncated]", "a plus") == \
        uk.normalize_quote_uk("a [truncated]", "a plus")
    # complete content -> no marker
    assert uk.normalize_quote_uk("повний текст", "повний текст") == \
        ("повний текст", "complete")
    # neither prefix nor equal -> explicit unverifiable state
    assert uk.normalize_quote_uk("чужий текст", "зовсім інший") == \
        (UNVERIFIABLE_QUOTE, "unverifiable")


def test_build_is_deterministic():
    wb1, analysis = _build()
    wb2 = openpyxl.Workbook()
    uk.build_all_uk(wb2, analysis)
    for title in uk.SHEET_NAMES_UK:
        assert _sheet_values(wb1, title) == _sheet_values(wb2, title)
