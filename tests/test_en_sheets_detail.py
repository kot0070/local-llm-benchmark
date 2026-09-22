"""Tests for tools/en_sheets_detail.py (Per-Model Test Detail + Failure Analysis)."""
import json
import os
import re
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import openpyxl  # noqa: E402

from tools.en_sheets_detail import (  # noqa: E402
    DETAIL_HEADER,
    FAILURE_HEADER,
    NO_PATTERN_ROW,
    NOT_AVAILABLE,
    SHEET_NAMES_DETAIL,
    build_detail_sheet,
    build_failures_sheet,
    load_analysis,
)

CYR = re.compile(r"[\u0400-\u04FF]")
BRITISH_STEMS = ("optimis", "behaviour", "organis", "centre",
                 "licence", "programme")
# US-correct words that contain the 'analys' stem; everything else with
# that stem (analyse, analyser, analysing, ...) is a British form.
ANALYS_ALLOW = {"analysis", "analyses", "analyst", "analysts", "analytic",
                "analytics", "analytical", "analytically"}


@pytest.fixture(scope="module")
def analysis():
    return load_analysis()


@pytest.fixture(scope="module")
def raw_json_text():
    chunks = []
    for name in ("group_a.json", "group_b.json", "group_c.json", "group_d.json"):
        with open(os.path.join(ROOT, "analysis_en", name), encoding="utf-8") as f:
            chunks.append(f.read())
    return "\n".join(chunks)


def _build(analysis):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_detail_sheet(wb, analysis)
    build_failures_sheet(wb, analysis)
    return wb


def british_hits(value):
    """British-stem word hits in a cell value ('analysis' etc. excluded)."""
    hits = []
    for word in re.findall(r"[A-Za-z]+", str(value or "")):
        wl = word.lower()
        if wl in ANALYS_ALLOW:
            continue
        if "analys" in wl or "analyz" in wl:
            # 'analyz' words are US forms (analyze, analyzing): not a hit.
            if "analys" in wl:
                hits.append(word)
            continue
        if any(st in wl for st in BRITISH_STEMS):
            hits.append(word)
    return hits


def test_detail_has_576_unique_rows_all_models(analysis):
    wb = openpyxl.Workbook()
    ws = build_detail_sheet(wb, analysis)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 576
    pairs = [(r[0], r[1]) for r in rows]
    assert len(set(pairs)) == 576
    assert len({r[0] for r in rows}) == 24
    # 24 contiguous per-model blocks, tests ascending inside each block.
    models = [r[0] for r in rows]
    assert models == sorted(models)
    for tag in sorted({r[0] for r in rows}):
        tids = [r[1] for r in rows if r[0] == tag]
        assert tids == sorted(tids), tag


def test_not_attempted_rows_empty_qsem_with_reason(analysis):
    wb = openpyxl.Workbook()
    ws = build_detail_sheet(wb, analysis)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    nas = [r for r in rows if r[5] == "Not attempted"]
    assert nas, "expected some not-attempted rows"
    for r in nas:
        assert r[3] is None, r[:2]  # Q_sem genuinely empty, never 0
        assert isinstance(r[6], str) and r[6].strip(), r[:2]
    for r in rows:
        assert r[5] in ("Strong", "Adequate", "Weak", "Failed",
                        "Not attempted", NOT_AVAILABLE), r[:2]


def test_not_attempted_rows_empty_cases_and_quality(analysis):
    wb = openpyxl.Workbook()
    ws = build_detail_sheet(wb, analysis)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    nas = [r for r in rows if r[5] == "Not attempted"]
    assert len(nas) == 361, len(nas)
    for r in nas:
        assert r[3] is None, r[:2]  # Q_sem genuinely empty, never 0
        assert r[4] is None, r[:2]  # Cases empty: no attempt count exists
    scored = [r for r in rows if r[5] != "Not attempted"]
    assert len(scored) == 215, len(scored)
    for r in scored:
        assert r[3] is not None, r[:2]
        assert r[4] is not None and str(r[4]).strip(), r[:2]


def test_failure_sheet_covers_every_model(analysis):
    wb = openpyxl.Workbook()
    ws = build_failures_sheet(wb, analysis)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    models = {m["model"] for m in analysis}
    assert {r[0] for r in rows} == models
    # sorted by model tag, then count descending within each model
    keys = [(r[0], -(r[2] if isinstance(r[2], int) else -1)) for r in rows]
    assert keys == sorted(keys)


def test_model_without_patterns_gets_placeholder_row():
    data = [{"model": "zzz-test:1b", "per_test": [],
             "failure_analysis": []}]
    wb = openpyxl.Workbook()
    ws = build_failures_sheet(wb, data)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][0] == "zzz-test:1b"
    assert rows[0][1] == NO_PATTERN_ROW
    assert all(v is None for v in rows[0][2:])


def test_quote_written_verbatim_unmodified():
    quote = "```json\n{\n \"answer\": \"1500 W\",\n \"citations\": [\"D1\"]\n}\n``` [...]"
    data = [{"model": "zzz-test:1b",
             "per_test": [{"test_id": "HOME-03", "title": "T",
                           "q_sem": 0.5, "n": "1/2",
                           "outcome": "adequate", "note": "N."}],
             "failure_analysis": [{"pattern": "FORMAT_ERROR[x]",
                                   "count": 3, "english_meaning": "M.",
                                   "what_the_model_did": "D.",
                                   "example_case_id": "C-1",
                                   "example_quote": quote,
                                   "why_it_matters": "W."}]}]
    wb = openpyxl.Workbook()
    ws = build_failures_sheet(wb, data)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][6] == quote
    ws2 = build_detail_sheet(openpyxl.Workbook(), data)
    assert list(ws2.iter_rows(min_row=2, values_only=True))[0][3] == 0.5


def test_missing_fields_render_not_available():
    data = [{"model": "zzz-test:1b",
             "per_test": [{"test_id": "HOME-01", "outcome": "superb"}],
             "failure_analysis": [{"pattern": "X", "count": "many"}]}]
    wb = openpyxl.Workbook()
    detail = list(build_detail_sheet(wb, data).iter_rows(
        min_row=2, values_only=True))[0]
    assert detail[2] == NOT_AVAILABLE  # title
    assert detail[3] == NOT_AVAILABLE  # q_sem unparseable
    assert detail[4] == NOT_AVAILABLE  # n
    assert detail[5] == NOT_AVAILABLE  # unknown outcome
    assert detail[6] == NOT_AVAILABLE  # note
    fail = list(build_failures_sheet(wb, data).iter_rows(
        min_row=2, values_only=True))[0]
    assert fail[2] == NOT_AVAILABLE  # non-int count
    assert fail[3] == NOT_AVAILABLE


def test_no_cyrillic_anywhere(analysis):
    wb = _build(analysis)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert not CYR.search(str(v or "")), (ws.title, v)


def test_british_spelling_only_flagged_passthrough(analysis, raw_json_text):
    # Module-authored fixed strings must be clean US English.
    for s in DETAIL_HEADER + FAILURE_HEADER + [NOT_AVAILABLE, NO_PATTERN_ROW,
                                               "Strong", "Adequate", "Weak",
                                               "Failed", "Not attempted"]:
        assert british_hits(s) == [], s
        assert not CYR.search(s), s
    # Every British-stem hit in a written cell must be verbatim passthrough
    # from the input JSON (flagged, not fixed -- see DONE report).
    wb = _build(analysis)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                for word in british_hits(v):
                    assert word in raw_json_text, (ws.title, word, v)


def test_deterministic(analysis):
    def values(wb):
        return [[c.value for c in row] for ws in wb.worksheets
                for row in ws.iter_rows()]
    assert values(_build(analysis)) == values(_build(analysis))
    assert SHEET_NAMES_DETAIL == ["Per-Model Test Detail", "Failure Analysis"]


def test_detail_formatting(analysis):
    wb = openpyxl.Workbook()
    ws = build_detail_sheet(wb, analysis)
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref is not None
    assert ws.column_dimensions["G"].width >= 60
    assert ws["D2"].number_format == "0.000"
