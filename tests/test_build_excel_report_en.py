"""Tests for tools/build_excel_report_en.py (no network, no Ollama)."""
import csv
import hashlib
import json
import os
import sys

import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from tools.build_excel_report_en import (  # noqa: E402
    EMPTY_OUTPUT_TEXT,
    NO_GEN_DATA,
    SHEET_NAMES,
    TRUNCATION_MARKER,
    UNVERIFIABLE_QUOTE,
    _is_not_ranked,
    americanize,
    british_hits,
    canonical_not_attempted,
    normalize_empty_quote_cells,
    normalize_evidence_quote,
    normalize_rating_basis,
    strip_trailing_marker,
    main,
)

DIGEST = "abc123"

EXPECTED_SHEETS = [
    "Executive Summary",
    "Run Overview",
    "Model Scorecard",
    "Recommendations",
    "Strengths & Weaknesses",
    "Per-Model Test Detail",
    "Failure Analysis",
    "HOME Verdicts",
    "Full Ranking",
    "Performance",
    "Known Model Behaviors",
    "HOME-07 Special Case",
]

RETIRED_SHEETS = ("Model Profiles", "Strengths & Hardware Fit")


def _rec(model, test, case, status, sem, sub=""):
    return {"key": "r|OLLAMA|%s@%s|%s@n1|%s|0|R0|p" % (model, DIGEST, test, case),
            "identity": {"tag": model, "digest": DIGEST,
                         "ollama_version": "0.34.2"},
            "response": {"content": "x", "thinking": "", "tool_calls": [],
                         "tool_call_channel": "none", "done_reason": "stop",
                         "prompt_eval_count": 1, "eval_count": 1},
            "timing": {"load_s": 0.1, "ttft_s": 0.2, "prompt_eval_s": 0.1,
                       "eval_s": 0.5, "prompt_tok_s": 10.0, "gen_tok_s": 2.0,
                       "wall_s": 1.0, "first_after_load": None},
            "resources": {"vram_peak_mb": 100.0, "vram_baseline_mb": 50.0,
                          "gpu_util_mean": 5.0, "temp_start": 40.0,
                          "temp_peak": 45.0, "throttle": False},
            "verdict": {"status": status, "sub_reason": sub,
                        "attribution": "MODEL", "sem": sem, "strict": sem,
                        "details": {}},
            "flags": [],
            "timestamps": {"utc": "2026-09-21T15:00:00+00:00",
                           "local": "2026-09-21T10:00:00-05:00"}}


def _write_run(run_dir, recs, with_perf=True, with_manifest=True):
    os.makedirs(run_dir, exist_ok=True)
    with open(os.path.join(run_dir, "results.jsonl"), "w",
              encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps(r) + "\n")
    if with_perf:
        with open(os.path.join(run_dir, "perf.csv"), "w", newline="",
                  encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["tag", "cold_load_s", "ttft_s", "prompt_tok_s",
                        "gen_tok_s", "vram_peak_mb", "offload_ratio",
                        "throttle"])
            w.writerow(["clean-ok:1b", 1.0, 0.2, 1000.0, 50.0, 500.0, 1.0,
                        False])
            w.writerow(["fail-model:7b", 9.0, 0.5, 800.0, 5.0, 7000, 0.5,
                        False])
    if with_manifest:
        man = {"run_id": os.path.basename(run_dir),
               "start_utc": "2026-09-21T14:00:00+00:00",
               "start_local": "2026-09-21T09:00:00-05:00",
               "budget_hours": 8,
               "ollama_version": "0.34.2",
               "fingerprint": {"gpu": "TEST-GPU", "cpu": "TEST-CPU",
                               "ram": "TEST-RAM", "os": "TEST-OS"}}
        with open(os.path.join(run_dir, "manifest.json"), "w",
                  encoding="utf-8") as f:
            json.dump(man, f)


def _synthetic_recs():
    recs = []
    for i in range(3):  # clean model, all OK
        recs.append(_rec("clean-ok:1b", "HOME-01", "HOME-01-c%d" % i,
                         "OK", 1.0))
    # failing model: 2 OK + 2 WRONG -> Q_sem 0.5 on HOME-03
    for i in range(2):
        recs.append(_rec("fail-model:7b", "HOME-03", "HOME-03-c%d" % i,
                         "OK", 1.0))
    for i in range(2, 4):
        recs.append(_rec("fail-model:7b", "HOME-03", "HOME-03-c%d" % i,
                         "WRONG_ANSWER", 0.0, sub="EVIDENCE_MISMATCH"))
    # UNSUPPORTED-only model (must show "Not ranked", no fabricated Q)
    for i in range(2):
        recs.append(_rec("unsup-model:1b", "HOME-02", "HOME-02-c%d" % i,
                         "UNSUPPORTED_CAPABILITY", 0.0, sub="LANG"))
    # NOT_RUN_BUDGET-only pair
    for i in range(2):
        recs.append(_rec("clean-ok:1b", "HOME-04", "HOME-04-c%d" % i,
                         "NOT_RUN_BUDGET", 0.0, sub="BUDGET"))
    return recs


def _open_xlsx(path):
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True)


def _non_quote_cells(wb):
    """Yield (sheet, coordinate, value) for every cell the US scan covers."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if ws.title == "Failure Analysis" and cell.column == 7:
                    continue
                if isinstance(cell.value, str) and cell.value:
                    yield ws.title, cell.coordinate, cell.value


def test_sheet_list_order_and_retired_absent(tmp_path):
    run_dir = str(tmp_path / "ORDER")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    assert wb.sheetnames == EXPECTED_SHEETS
    assert SHEET_NAMES == EXPECTED_SHEETS
    assert len(SHEET_NAMES) == 12
    for retired in RETIRED_SHEETS:
        assert retired not in wb.sheetnames


def test_synthetic_sheets_values_and_determinism(tmp_path):
    from bench import report as rep
    run_dir = str(tmp_path / "SYNTH")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    out = os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx")
    assert os.path.exists(out)
    with open(out, "rb") as f:
        first_bytes = f.read()
    assert main([run_dir]) == 0  # rerun: byte-identical output
    with open(out, "rb") as f:
        assert f.read() == first_bytes
    wb = _open_xlsx(out)
    assert wb.sheetnames == SHEET_NAMES
    assert len(SHEET_NAMES) == 12
    # header row present and bold on every sheet
    for ws in wb.worksheets:
        assert ws.max_row >= 1 and ws.max_column >= 1
        assert ws["A1"].value not in (None, "")
        assert ws["A1"].font.bold is True
        assert ws.freeze_panes is not None
    # ranking: fail-model HOME-03 Q_sem matches bench.report independently
    records = rep.load_records(run_dir)
    per_test, _, _ = rep.compute_tables(records)
    exp = next(r for r in per_test["HOME-03"]
               if r["model"] == "fail-model:7b")
    ws = wb["Full Ranking"]
    found = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == "HOME-03" and row[1] == "fail-model:7b":
            found = row
            break
    assert found is not None
    assert found[2] == pytest.approx(exp["q_sem"], abs=1e-9)
    # UNSUPPORTED-only row labelled, not ranked with a fabricated score
    ws_rows = list(ws.iter_rows(min_row=2, values_only=True))
    unsup = [r for r in ws_rows
             if r[0] == "HOME-02" and r[1] == "unsup-model:1b"]
    assert len(unsup) == 1
    assert unsup[0][-1] == "Not ranked (unsupported / time-boxed out)"
    # NOT_RUN_BUDGET-only row labelled as well
    notrun = [r for r in ws_rows
              if r[0] == "HOME-04" and r[1] == "clean-ok:1b"]
    assert len(notrun) == 1
    assert notrun[0][-1] == "Not ranked (unsupported / time-boxed out)"
    # verdicts sheet: one row per HOME test (24) + header, never blank
    wv = wb["HOME Verdicts"]
    assert wv.max_row == 25
    for row in wv.iter_rows(min_row=2, values_only=True):
        assert row[-1] not in (None, "")
    # HOME-07 sheet exists with the non-comparability note
    w7 = wb["HOME-07 Special Case"]
    texts = " ".join(str(c.value or "") for row in w7.iter_rows()
                     for c in row)
    assert "cannot be compared like-for-like" in texts
    # executive summary exists with real prose sections + per-model pointer
    we = wb["Executive Summary"]
    assert we.max_row >= 7
    assert we["A1"].value == "Section"
    sections = [r[0] for r in we.iter_rows(min_row=2, values_only=True)]
    assert "Per-model depth" in sections
    # known behaviors sheet is English prose, one row per entry
    wk = wb["Known Model Behaviors"]
    assert wk.max_row >= 2


def test_missing_gen_tok_s_writes_no_data(tmp_path):
    run_dir = str(tmp_path / "NODATA")
    recs = [_rec("embed-only:1b", "HOME-02", "HOME-02-c0",
                  "UNSUPPORTED_CAPABILITY", 0.0, sub="LANG")]
    for i in range(2):
        recs.append(_rec("text-model:1b", "HOME-01", "HOME-01-c%d" % i,
                          "OK", 1.0))
    os.makedirs(run_dir, exist_ok=True)
    with open(os.path.join(run_dir, "results.jsonl"), "w",
              encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps(r) + "\n")
    with open(os.path.join(run_dir, "perf.csv"), "w", newline="",
              encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["tag", "cold_load_s", "ttft_s", "prompt_tok_s",
                    "gen_tok_s", "vram_peak_mb", "offload_ratio",
                    "throttle"])
        w.writerow(["embed-only:1b", 1.0, 0.2, 1000.0, "", 500.0, 1.0,
                    False])
        w.writerow(["text-model:1b", 1.0, 0.2, 1000.0, 50.0, 500.0, 1.0,
                    False])
    man = {"run_id": "NODATA",
           "start_local": "2026-09-21T09:00:00-05:00",
           "budget_hours": 8}
    with open(os.path.join(run_dir, "manifest.json"), "w",
              encoding="utf-8") as f:
        json.dump(man, f)
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    spd = wb["Performance"]
    spd_gen = next(c.column for c in spd[1]
                   if str(c.value or "") == "Gen tok/s")
    emb_spd = next(r for r in spd.iter_rows(min_row=2, values_only=True)
                   if r[0] == "embed-only:1b")
    assert emb_spd[spd_gen - 1] == NO_GEN_DATA
    txt_spd = next(r for r in spd.iter_rows(min_row=2, values_only=True)
                   if r[0] == "text-model:1b")
    assert txt_spd[spd_gen - 1] == 50.0


def test_empty_results_no_crash(tmp_path):
    run_dir = str(tmp_path / "EMPTY")
    _write_run(run_dir, [], with_perf=False, with_manifest=False)
    assert main([run_dir]) == 0
    out = os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx")
    assert os.path.exists(out)
    wb = _open_xlsx(out)
    assert wb.sheetnames == SHEET_NAMES


REAL_RUN = os.path.join(ROOT, "results", "night_20260921-155146")


def _all_texts(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v:
                    yield v


def test_no_cyrillic_anywhere_synthetic(tmp_path):
    import re
    run_dir = str(tmp_path / "CYR")
    recs = _synthetic_recs()
    recs.append(_rec("cyr-model:1b", "HOME-03", "HOME-03-c9",
                     "UNSUPPORTED_CAPABILITY", 0.0, sub="поза scope"))
    recs.append(_rec("cyr-model:1b", "HOME-05", "HOME-05-c9",
                     "UNSUPPORTED_CAPABILITY", 0.0, sub="CAP: немає vision"))
    _write_run(run_dir, recs)
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    cyr = re.compile(r"[\u0400-\u04FF]")
    for t in _all_texts(wb):
        assert not cyr.search(t), t[:120]
        assert "поза" not in t
        assert "немає" not in t


def test_no_cyrillic_real_run():
    import re
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    import openpyxl
    wb = openpyxl.load_workbook(out, data_only=True)
    cyr = re.compile(r"[\u0400-\u04FF]")
    for t in _all_texts(wb):
        assert not cyr.search(t), t[:160]


def test_no_british_spelling_stem_scan(tmp_path):
    run_dir = str(tmp_path / "BRIT")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    assert "Known Model Behaviors" in wb.sheetnames
    # Stems without a trailing boundary: catches optimisation, not just
    # optimise (the previous round's miss).
    assert british_hits("optimisation")
    assert british_hits("Optimisation")
    assert british_hits("behaviours")
    assert british_hits("normalised")
    assert british_hits("Constrained optimisation, JSON solution")
    assert british_hits("travelling")
    assert british_hits("modelled")
    assert british_hits("analysing")
    # US spellings pass clean, including analysis/analyses/analyst.
    assert british_hits("optimization") == []
    assert british_hits("analysis of the results") == []
    assert british_hits("organize the schedule") == []
    for _title, _coord, text in _non_quote_cells(wb):
        assert british_hits(text) == [], (_title, _coord, text[:120])
    # The previously shipped defect is gone at its known address.
    assert wb["HOME Verdicts"]["B18"].value == \
        "Constrained optimization, JSON solution"


def test_americanize_normalizes_synthetic_sheet():
    import openpyxl
    from tools.build_excel_report_en import _americanize_workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HOME Verdicts"
    ws.append(["Title"])
    ws.append(["Constrained optimisation, JSON solution"])
    ws.append(["Optimisation of favours and behaviours"])
    ws.append(["The model travelled and modelled labelling"])
    ws.append(["Analysis complete, center licenced defence"])
    _americanize_workbook(wb)
    assert ws["A2"].value == "Constrained optimization, JSON solution"
    assert ws["A3"].value == "Optimization of favors and behaviors"
    assert ws["A4"].value == "The model traveled and modeled labeling"
    assert ws["A5"].value == "Analysis complete, center licensed defense"
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert british_hits(row[0]) == [], row[0]
    # Verbatim quotes are exempt: the quote column keeps British text.
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Failure Analysis"
    for _ in range(6):
        ws2.append(["x"] * 8)
    ws2.cell(row=2, column=7, value="the state was cancelled by optimisation")
    _americanize_workbook(wb2)
    assert ws2.cell(row=2, column=7).value == \
        "the state was cancelled by optimisation"


def test_americanize_units():
    assert americanize("optimisation") == "optimization"
    assert americanize("Optimisation") == "Optimization"
    assert americanize("behaviours") == "behaviors"
    assert americanize("favours rapid work") == "favors rapid work"
    assert americanize("normalised") == "normalized"
    assert americanize("the data center license") == \
        "the data center license"
    # analyse family converts; analysis/analyses/analyst never match.
    assert americanize("analysing the logs") == "analyzing the logs"
    assert americanize("analysis of the results") == "analysis of the results"
    # L1 rule-based cases: the words that twice shipped British.
    assert americanize("localisation") == "localization"
    assert americanize("Localisation output") == "Localization output"
    assert americanize("optimisation") == "optimization"
    assert americanize("behaviours") == "behaviors"
    assert americanize("travelled") == "traveled"
    assert americanize("modelling the labels") == "modeling the labels"
    assert americanize("organised favourably") == "organized favorably"
    assert americanize("CENTRES DEFENCES") == "CENTERS DEFENSES"
    # Negative cases: genuine US words the rules must leave alone.
    for w in ("surprise", "surprises", "analysis", "analyses", "exercise",
              "exercised", "enterprise", "enterprises", "promise",
              "precise", "otherwise", "this", "crisis", "glamour"):
        assert americanize(w) == w, w
        assert british_hits(w) == [], w
    for w in ("localisation", "optimisation", "behaviours", "travelled",
              "normalised"):
        assert british_hits(w) != [], w


def test_quote_helpers_units():
    assert strip_trailing_marker("abc [truncated]") == "abc"
    assert strip_trailing_marker("abc[...]") == "abc"
    assert strip_trailing_marker("abc [...]  ") == "abc"
    assert strip_trailing_marker("abc...") == "abc"
    assert strip_trailing_marker("plain quote") == "plain quote"
    rendered, status = normalize_evidence_quote("abc", "abcdef")
    assert (rendered, status) == ("abc" + TRUNCATION_MARKER, "truncated")
    rendered, status = normalize_evidence_quote("abc [truncated]", "abcdef")
    assert (rendered, status) == ("abc" + TRUNCATION_MARKER, "truncated")
    rendered, status = normalize_evidence_quote("abc", "abc")
    assert (rendered, status) == ("abc", "complete")
    rendered, status = normalize_evidence_quote("xyz", "abcdef")
    assert (rendered, status) == (UNVERIFIABLE_QUOTE, "unverifiable")
    rendered, status = normalize_evidence_quote("abc", None)
    assert (rendered, status) == (UNVERIFIABLE_QUOTE, "unverifiable")


def test_canonical_not_attempted_units():
    elig = {
        "HOME-05": {"m:1b": {"code": "U",
                             "reason": "CAP: немає vision"}},
        "HOME-01": {"m:1b": {"code": "S", "reason": "поза scope"}},
        "HOME-08": {"m:1b": {"code": "O", "reason": ""}},
        "HOME-18": {"m:1b": {"code": "O",
                             "reason": "багатоходовий діалог не задокументовано"}},
        "HOME-09": {"m:1b": {"code": "U",
                             "reason": "CONTEXT_WINDOW: потрібно 36000, runtime 8192"}},
    }
    # Identical causes render identically whatever the original phrasing.
    for note in ("Never attempted: the model has no documented vision "
                 "capability.",
                 "Not attempted: missing documented capability, no vision "
                 "support.",
                 "Never attempted: the model is not documented for vision "
                 "input."):
        assert canonical_not_attempted("HOME-05", "m:1b", note, elig) == \
            "Not attempted — CAP: missing vision modality."
    assert canonical_not_attempted("HOME-01", "m:1b",
                                   "Not attempted: outside the planned scope "
                                   "of this run for this model.", elig) == \
        "Not attempted — outside the planned scope of this run for this model."
    assert canonical_not_attempted("HOME-08", "m:1b",
                                   "Never attempted: the pair was not "
                                   "scheduled in this run and no reason was "
                                   "recorded.", elig) == \
        "Not attempted — the pair was not scheduled in this run."
    assert canonical_not_attempted("HOME-18", "m:1b",
                                   "Never attempted: multi-turn dialogue is "
                                   "not a documented scenario for this "
                                   "model.", elig) == \
        "Not attempted — multi-turn dialogue not documented."
    assert canonical_not_attempted("HOME-09", "m:1b",
                                   "Never attempted: the test needs about "
                                   "36000 tokens of context and the runtime "
                                   "allows 8192.", elig) == \
        ("Not attempted — CONTEXT_WINDOW: requires 36000 tokens, "
         "runtime 8192 tokens.")
    # Genuinely model-specific detail survives after the canonical clause.
    got = canonical_not_attempted("HOME-99", "m:1b",
                                  "Not attempted: the 8-hour run budget was "
                                  "exhausted before this block.", {})
    assert got.startswith("Not attempted — no scored cases recorded in "
                          "this run. ")
    assert "budget was exhausted" in got
    # Retired prefix never appears.
    for got2 in (canonical_not_attempted("HOME-05", "m:1b", "x", elig), got):
        assert "Never attempted" not in got2


def test_glossed_sub_reason_codes(tmp_path):
    from tools.build_excel_report_en import _gloss_sub_reason
    assert _gloss_sub_reason("поза scope") == "out-of-scope test"
    assert _gloss_sub_reason("CAP: немає text") == "CAP: missing text modality"
    assert _gloss_sub_reason("CAP: немає vision") == \
        "CAP: missing vision modality"
    assert _gloss_sub_reason("-") == ""
    assert "потрібно" not in _gloss_sub_reason(
        "CONTEXT_WINDOW: потрібно 13500, runtime 8192")
    run_dir = str(tmp_path / "GLOSS")
    recs = [_rec("m:1b", "HOME-03", "HOME-03-c0",
                 "UNSUPPORTED_CAPABILITY", 0.0, sub="поза scope")]
    _write_run(run_dir, recs)
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    ws = wb["Per-Model Test Detail"]
    texts = " ".join(str(c.value or "") for row in ws.iter_rows()
                     for c in row)
    # Canonical notes reuse the gloss machinery: no Cyrillic, no [-].
    assert "CAP: missing vision modality" in texts
    assert "[-]" not in texts


def test_behaviors_rows_have_subjects(tmp_path):
    run_dir = str(tmp_path / "BEH")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    ws = wb["Known Model Behaviors"]
    assert ws["B1"].value == "Observed behavior"
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 12
    assert rows[1][0] == "granite-code:8b-instruct"
    assert rows[10][0] == "glm-ocr:latest"
    for _model, text in rows:
        assert text and text[0].isupper(), text[:80]
        assert " px" not in text and "px)" not in text
    blob = " ".join(t for _m, t in rows)
    assert "16K-token" in blob and "8K-token" in blob


def test_home07_sheet_structure(tmp_path):
    run_dir = str(tmp_path / "H07")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    ws = wb["HOME-07 Special Case"]
    assert ws["A4"].value == "Cases (scored/total)"
    assert ws["A5"].value == "Run"
    assert ws["B5"].value == "Main run"
    assert ws["C5"].value == "Follow-up run"
    assert ws["A6"].value == "Outcome counts"
    by_label = {r[0]: r for r in ws.iter_rows(min_row=1, values_only=True)
                if r[0]}
    assert "What was tried" in by_label
    assert "NOT_RUN_BUDGET placeholder" in str(by_label["What was tried"][2] or "")
    assert "cannot be compared like-for-like" in str(by_label["Why it matters"][1] or "")


def test_performance_headers_and_formats(tmp_path):
    run_dir = str(tmp_path / "PERF")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    ws = wb["Performance"]
    headers = [c.value for c in ws[1]]
    assert headers[2] == "Time to first token (s)"
    assert headers[5] == "Offload ratio (1.0 = fully on GPU)"
    assert ws["B2"].number_format == "0.0"
    assert ws["E2"].number_format == "0"


def test_verdict_labels_parallel(tmp_path):
    run_dir = str(tmp_path / "VERD")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    wb = _open_xlsx(os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx"))
    ws = wb["HOME Verdicts"]
    verdicts = [r[9] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert "Tie" not in verdicts
    assert any(v == "Home model ties" for v in verdicts) or \
        all(v == "Insufficient data for a verdict" for v in verdicts)
    comps = [r[5] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert not any(c == "Insufficient data for a verdict" for c in comps)


def test_detail_sheet_shape_and_canonical_notes():
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    assert "Model Scorecard" in wb.sheetnames
    assert "Recommendations" in wb.sheetnames
    assert "Strengths & Weaknesses" in wb.sheetnames
    ws = wb["Per-Model Test Detail"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 576
    assert len({(r[0], r[1]) for r in rows}) == 576
    assert len({r[0] for r in rows}) == 24
    nas = [r for r in rows if r[5] == "Not attempted"]
    assert len(nas) == 361
    for r in nas:
        assert r[3] is None, r[:2]  # Q_sem genuinely empty, never 0
        assert isinstance(r[6], str) and r[6].startswith("Not attempted — "), \
            r[:2]
    for r in rows:
        note = r[6] if isinstance(r[6], str) else ""
        assert "Never attempted" not in note, r[:2]


def test_failure_quotes_verified_real_run():
    from bench import report as rep
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ws = wb["Failure Analysis"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 133
    records = rep.load_records(REAL_RUN)
    index = {}
    for r in records:
        tag = rep.model_of(r)
        case = str(r.get("key", "")).split("|")[4]
        index.setdefault((tag, case),
                         (r.get("response") or {}).get("content", ""))
    complete = truncated = unverifiable = quoted = 0
    for r in rows:
        q = r[6]
        if q in (None, "Not available", EMPTY_OUTPUT_TEXT) or \
                q == "No failure pattern reached the reporting threshold " \
                     "in this run":
            continue
        quoted += 1
        assert q != UNVERIFIABLE_QUOTE, (r[0], r[5])
        assert "[...]" not in q, (r[0], r[5])
        stripped = strip_trailing_marker(q)
        assert "[truncated]" not in stripped, (r[0], r[5])
        full = index.get((r[0], r[5]))
        assert full is not None, (r[0], r[5])
        if stripped == full:
            complete += 1
        else:
            assert full.startswith(stripped) and stripped, (r[0], r[5])
            assert q.endswith(TRUNCATION_MARKER), (r[0], r[5])
            truncated += 1
    assert quoted == 110
    assert complete == 53
    assert truncated == 57
    assert unverifiable == 0


def test_spotcheck_audited_numbers_unchanged():
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ov = {r[0]: r[1] for r in wb["Run Overview"].iter_rows(values_only=True)
          if r[0]}
    assert ov["run_id"] == "night_20260921-155146"
    assert ov["Models run"] == 24
    assert ov["Test groups run"] == 25
    assert ov["Records"] == 4190
    assert ov["NOT_RUN_BUDGET records"] == 799
    assert ov["NOT_RUN_BUDGET share (%)"] == 19.1
    assert ov["Duration (min)"] == 461.5
    hv = {r[0]: r for r in wb["HOME Verdicts"].iter_rows(min_row=2,
                                                        values_only=True)}
    assert hv["HOME-01"][3] == 0.858  # home Q_sem
    assert hv["HOME-01"][6] == 0.944  # best competitor Q_sem
    assert hv["HOME-03"][3] == 0  # command-r7b home zero
    fr = list(wb["Full Ranking"].iter_rows(min_row=2, values_only=True))
    assert len(fr) == 576
    top = next(r for r in fr if r[0] == "HOME-01" and r[1] == "qwen3:8b")
    assert (top[2], top[3]) == (1, 1)
    pf = {r[0]: r for r in wb["Performance"].iter_rows(min_row=2,
                                                      values_only=True)
          if r[0]}
    assert pf["functiongemma:270m"][4] == pytest.approx(460.974, abs=0.01)
    h7 = list(wb["HOME-07 Special Case"].iter_rows(values_only=True))
    assert (h7[1][1], h7[1][2]) == (0, 0)
    assert (h7[3][1], h7[3][2]) == ("16/16", "16/17")
    ex = {r[0]: r[1] for r in wb["Executive Summary"].iter_rows(min_row=2,
                                                               values_only=True)}
    assert "0 wins" in ex["Home-model verdicts"]


def test_determinism_sha(tmp_path):
    run_dir = str(tmp_path / "DET")
    _write_run(run_dir, _synthetic_recs())
    assert main([run_dir]) == 0
    out = os.path.join(run_dir, "NIGHT1_REPORT_EN.xlsx")
    with open(out, "rb") as f:
        first = f.read()
    assert main([run_dir]) == 0
    with open(out, "rb") as f:
        second = f.read()
    assert first == second
    assert hashlib.sha256(first).hexdigest() == \
        hashlib.sha256(second).hexdigest()


def test_real_run_exit_0_and_readonly():
    if not os.path.isdir(REAL_RUN):
        pytest.skip("missing real run dir")
    before = set(os.listdir(REAL_RUN))
    assert main(["night_20260921-155146"]) == 0
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    assert os.path.exists(out)
    wb = _open_xlsx(out)
    assert wb.sheetnames == SHEET_NAMES
    # tool writes only its own declared output file
    new_files = set(os.listdir(REAL_RUN)) - before
    assert new_files <= {"NIGHT1_REPORT_EN.xlsx"}, new_files


def _prose_rating_mentions(obj, skip_keys, out, path=""):
    import re
    rx = re.compile(r"(\d+\.\d+)\s+ratings?"
                    r"|ratings?\s+(?:of\s+|at\s+|is\s+)?"
                    r"(?:a\s+|an\s+)?(\d+\.\d+)", re.IGNORECASE)
    if isinstance(obj, dict):
        for k, v in obj.items():
            _prose_rating_mentions(v, skip_keys, out,
                                   path + "/" + str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            _prose_rating_mentions(v, skip_keys, out,
                                   path + "/" + str(i))
    elif isinstance(obj, str):
        if path.rsplit("/", 1)[-1] in skip_keys:
            return
        for m in rx.finditer(obj):
            out.append((path, m.group(1) or m.group(2)))


def test_prose_ratings_match_rating_column():
    # N1 class test: no prose field may quote a rating that disagrees
    # with the model's own rating_10 (stale-sentence detector).
    import glob
    files = sorted(glob.glob(os.path.join(ROOT, "analysis_en",
                                          "group_*.json")))
    assert len(files) == 4
    checked = 0
    for path in files:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        for model in data:
            rating = model.get("overall", {}).get("rating_10")
            assert isinstance(rating, (int, float)), model.get("model")
            expected = "%.1f" % rating
            found: list = []
            _prose_rating_mentions(
                model, {"model", "test_id", "example_case_id"}, found)
            # No requirement that every model mentions its rating in
            # prose (most do not); every mention that exists must agree.
            for where, num in found:
                assert num == expected, (model.get("model"), where,
                                         num, expected)
                checked += 1
    assert checked >= 10


def test_no_ranked_with_zero_scored():
    # P1 class test: the ranking predicate keys on scored cases, not on
    # whichever status mix produced the zero.
    overflow = {"counts": {"CONTEXT_OVERFLOW": 2, "NOT_RUN_BUDGET": 1},
                "n_scored": 0, "n_total": 3}
    assert _is_not_ranked(overflow) is True
    assert _is_not_ranked({"counts": {"NOT_RUN_BUDGET": 2},
                           "n_scored": 0}) is True
    assert _is_not_ranked({"counts": {"UNSUPPORTED_CAPABILITY": 2},
                           "n_scored": 0}) is True
    assert _is_not_ranked({"counts": {"OK": 3}, "n_scored": 3}) is False
    assert _is_not_ranked({"counts": {"OK": 1, "WRONG_ANSWER": 2},
                           "n_scored": 3}) is False
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ws = wb["Full Ranking"]
    header = [c.value for c in ws[1]]
    n_scored_i = header.index("N scored")
    status_i = header.index("Ranking status")
    ranked = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[status_i] == "Ranked":
            assert row[n_scored_i] not in (0, None), row[:2]
            ranked += 1
    assert ranked == 215


def test_rating_basis_canonical_units():
    models = [
        {"model": "m1",
         "overall": {"rating_basis":
                     "Mean semantic quality 0.855 across 6 scored tests, "
                     "home-test quality 0.858, and an OK rate of 30.7 "
                     "percent combine to a rating of 7.5."}},
        {"model": "m2",
         "overall": {"rating_basis":
                     "Mean semantic quality 0.945 over 1 scored test, "
                     "home-test quality 0.945, and an OK rate of 0.188. "
                     "The rating rests on a single test."}},
        {"model": "m3",
         "overall": {"rating_basis":
                     "Mean semantic quality 0.980 on its single scored "
                     "test HOME-11, own-test quality 0.980, and an "
                     "exact-match rate of 0.900. The rating rests on one "
                     "test only."}},
        {"model": "m4",
         "overall": {"rating_basis":
                     "Mean quality 0.697 across 21 scored tests, home-test "
                     "quality 0.900, exact-answer rate 0.670. The rating "
                     "follows the shared formula."}},
    ]
    assert normalize_rating_basis(models) == {"normalized": 4,
                                             "unparsed": 0}
    got = [m["overall"]["rating_basis"] for m in models]
    assert got[0] == ("Mean semantic quality 0.855 across 6 scored tests; "
                      "home-test quality 0.858; OK rate 0.307.")
    assert got[1] == ("Mean semantic quality 0.945 across 1 scored test; "
                      "home-test quality 0.945; OK rate 0.188. "
                      "The rating rests on a single test.")
    assert got[2] == ("Mean semantic quality 0.980 across 1 scored test "
                      "(HOME-11); home-test quality 0.980; OK rate 0.900. "
                      "The rating rests on one test only.")
    assert got[3] == ("Mean semantic quality 0.697 across 21 scored tests; "
                      "home-test quality 0.900; OK rate 0.670. The rating "
                      "follows the shared formula.")
    # Unknown shapes are left untouched, never mangled.
    odd = [{"model": "mx", "overall": {"rating_basis": "hand-written"}}]
    assert normalize_rating_basis(odd) == {"normalized": 0, "unparsed": 1}
    assert odd[0]["overall"]["rating_basis"] == "hand-written"


def test_rating_basis_canonical_real_run():
    import re
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ws = wb["Model Scorecard"]
    header = [c.value for c in ws[3]]
    basis_i = header.index("Rating basis")
    pat = re.compile(
        r"^Mean semantic quality \d+\.\d+ across \d+ scored tests?"
        r"( \(HOME-\d+\))?; home-test quality \d+\.\d+( on HOME-\d+)?; "
        r"OK rate \d+\.\d+")
    bases = [r[basis_i] for r in ws.iter_rows(min_row=4, values_only=True)
             if isinstance(r[basis_i], str) and r[basis_i]]
    assert len(bases) == 24
    for basis in bases:
        assert pat.match(basis), basis
        assert "percent" not in basis


def test_empty_quote_cells_units():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Failure Analysis"
    ws.append(["Model", "Pattern", "Count", "Meaning", "What", "Case",
               "Example output", "Why"])
    ws.append(["m:1b", "p", 1, "x", "y", "c", "Not available", "z"])
    ws.append(["m:1b", "p", 1, "x", "y", "c", "real quote", "z"])
    assert normalize_empty_quote_cells(wb) == 1
    assert ws.cell(row=2, column=7).value == EMPTY_OUTPUT_TEXT == \
        "No output \u2014 the model produced no content for this case"
    assert ws.cell(row=3, column=7).value == "real quote"


def test_empty_quote_cells_real_run():
    out = os.path.join(REAL_RUN, "NIGHT1_REPORT_EN.xlsx")
    if not os.path.exists(out):
        pytest.skip("missing real workbook")
    wb = _open_xlsx(out)
    ws = wb["Failure Analysis"]
    header = [c.value for c in ws[1]]
    quote_i = header.index("Example output")
    empties = [r for r in ws.iter_rows(min_row=2, values_only=True)
               if r[quote_i] == EMPTY_OUTPUT_TEXT]
    assert len(empties) == 23
    for r in ws.iter_rows(values_only=True):
        for v in r:
            assert v != "Not available", r[:2]
