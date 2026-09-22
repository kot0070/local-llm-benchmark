"""Per-model detail sheets for the English NIGHT-1 workbook. Stdlib + openpyxl only.

Renders two sheets from the already-audited deep-analysis JSON
(analysis_en/group_{a,b,c,d}.json):

  1. "Per-Model Test Detail" -- one row per (model, test) pair (576 rows),
     sourced from each model's ``per_test[]`` entries.
  2. "Failure Analysis" -- one row per failure pattern per model, sourced
     from each model's ``failure_analysis[]`` entries.

This module presents only what the JSON says: no re-derivation from raw
records, no new claims about any model. A missing field renders as the
explicit string ``Not available``. Shared formatting helpers (header
style, quality color scale) are reused from tools/build_excel_report_en.py
(read-only import). Never contacts Ollama.

Determinism: rows are emitted in a fixed sort order (model tag ascending,
then test id ascending / count descending) with fixed column widths, so
two builds over the same input produce identical cell values.
"""
from __future__ import annotations

import json
import os

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
ANALYSIS_DIR = os.path.join(ROOT, "analysis_en")
GROUP_FILES = ("group_a.json", "group_b.json", "group_c.json", "group_d.json")

SHEET_NAMES_DETAIL = ["Per-Model Test Detail", "Failure Analysis"]

DETAIL_HEADER = ["Model", "Test", "Test title", "Q_sem",
                 "Cases (scored/total)", "Outcome", "Notes"]
FAILURE_HEADER = ["Model", "Failure pattern", "Count", "What it means",
                  "What the model did", "Example case", "Example output",
                  "Why it matters"]

NOT_AVAILABLE = "Not available"
NO_PATTERN_ROW = ("No failure pattern reached the "
                  "reporting threshold in this run")

OUTCOME_LABELS = {
    "strong": "Strong",
    "adequate": "Adequate",
    "weak": "Weak",
    "failed": "Failed",
    "not attempted": "Not attempted",
}


def load_analysis(root: str | None = None) -> list[dict]:
    """Load and merge group_{a,b,c,d}.json, sorted by model tag ascending.

    Defined locally because tools/en_sheets_narrative.py (which will own
    the canonical loader) does not exist yet; the integrator reconciles
    the duplicate. Behavior contract: merged list of all model objects,
    sorted by ``model`` tag.
    """
    base = os.path.join(os.path.abspath(root), "analysis_en") if root else ANALYSIS_DIR
    models: list[dict] = []
    for name in GROUP_FILES:
        with open(os.path.join(base, name), encoding="utf-8") as f:
            models.extend(json.load(f))
    models.sort(key=lambda m: str(m.get("model", "")))
    return models


def _outcome_label(raw: object) -> str:
    return OUTCOME_LABELS.get(str(raw or "").strip().lower(), NOT_AVAILABLE)


def _text(value: object) -> str:
    s = str(value).strip() if value is not None else ""
    return s if s else NOT_AVAILABLE


def build_detail_sheet(wb, analysis: list[dict]):
    """Append the 'Per-Model Test Detail' sheet (one row per model x test)."""
    from openpyxl.styles import Alignment, Border, Side
    from tools.build_excel_report_en import _q_scale, _style_table

    ws = wb.create_sheet(SHEET_NAMES_DETAIL[0])
    ws.append(DETAIL_HEADER)
    first_rows: set[int] = set()
    for model in analysis:
        tag = str(model.get("model", "") or "").strip() or NOT_AVAILABLE
        entries = sorted(model.get("per_test") or [],
                         key=lambda e: str(e.get("test_id", "")))
        first_rows.add(ws.max_row + 1)
        for entry in entries:
            outcome_raw = str(entry.get("outcome", "") or "").strip().lower()
            not_attempted = outcome_raw == "not attempted"
            q_raw = entry.get("q_sem")
            if not_attempted:
                q_cell: object = None  # genuinely empty: no score exists
            elif isinstance(q_raw, bool) or q_raw is None or (
                    isinstance(q_raw, str) and not q_raw.strip()):
                q_cell = NOT_AVAILABLE
            else:
                try:
                    q_cell = round(float(q_raw), 3)
                except (TypeError, ValueError):
                    q_cell = NOT_AVAILABLE
            ws.append([tag,
                       _text(entry.get("test_id")),
                       _text(entry.get("title")),
                       q_cell,
                       None if not_attempted else _text(entry.get("n")),
                       _outcome_label(entry.get("outcome")),
                       _text(entry.get("note"))])
    _style_table(ws, [22, 10, 40, 10, 18, 14, 72], {3: "0.000"})
    ws.freeze_panes = "B2"  # keep the header row and the Model column visible
    ws.auto_filter.ref = ws.dimensions
    _q_scale(ws, "D", ws.max_row)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    block_top = Border(top=Side(style="thin", color="BFBFBF"))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(DETAIL_HEADER)):
        for cell in row:
            cell.alignment = wrap_top
            if cell.row in first_rows:
                cell.border = block_top
        ws.row_dimensions[row[0].row].height = 30
    return ws


def build_failures_sheet(wb, analysis: list[dict]):
    """Append the 'Failure Analysis' sheet (one row per pattern per model)."""
    from openpyxl.styles import Alignment, Font
    from tools.build_excel_report_en import _style_table

    ws = wb.create_sheet(SHEET_NAMES_DETAIL[1])
    ws.append(FAILURE_HEADER)
    quote_font = Font(name="Consolas")
    for model in analysis:
        tag = str(model.get("model", "") or "").strip() or NOT_AVAILABLE
        patterns = sorted(model.get("failure_analysis") or [],
                          key=lambda e: (-(e.get("count")
                                           if isinstance(e.get("count"), int)
                                           and not isinstance(e.get("count"), bool)
                                           else -1),
                                         str(e.get("pattern", ""))))
        if not patterns:
            ws.append([tag, NO_PATTERN_ROW, None, None, None, None, None, None])
            continue
        for entry in patterns:
            count = entry.get("count")
            count_cell: object = (count if isinstance(count, int)
                                  and not isinstance(count, bool)
                                  else NOT_AVAILABLE)
            ws.append([tag,
                       _text(entry.get("pattern")),
                       count_cell,
                       _text(entry.get("english_meaning")),
                       _text(entry.get("what_the_model_did")),
                       _text(entry.get("example_case_id")),
                       entry.get("example_quote")
                       if isinstance(entry.get("example_quote"), str)
                       and entry.get("example_quote", "").strip()
                       else NOT_AVAILABLE,
                       _text(entry.get("why_it_matters"))])
    _style_table(ws, [22, 30, 8, 42, 52, 14, 52, 42])
    ws.auto_filter.ref = ws.dimensions
    wrap_top = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(FAILURE_HEADER)):
        for cell in row:
            cell.alignment = wrap_top
        ws.row_dimensions[row[0].row].height = 60
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7,
                            max_row=ws.max_row):
        for cell in row:
            if isinstance(cell.value, str) and cell.value != NOT_AVAILABLE:
                cell.font = quote_font  # verbatim model output: monospace
    return ws
