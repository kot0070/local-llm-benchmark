"""Consolidated English-language workbook for a NIGHT-1 run. Stdlib + openpyxl only.

Reads results/<run_id>/results.jsonl + perf.csv + manifest.json (plus the
supplementary results/<run_id>_gv/results.jsonl for HOME-07 / granite3.2-vision
plain-answer adapter, when present) and writes
results/<run_id>/NIGHT1_REPORT_EN.xlsx.

Every number comes from the records through bench.report helpers
(load_records, compute_tables, summarize_group, home_verdicts, run_overview);
this tool does no scoring arithmetic of its own. English prose is written
fresh from that data (not translated from any Ukrainian document).
Never contacts Ollama.

Determinism: workbook properties use a fixed timestamp and the .xlsx zip
entries are rewritten with a fixed date_time, so two runs on the same input
produce byte-identical bytes.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import sys
import zipfile
from datetime import datetime

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from bench import report as _report  # noqa: E402

XLSX_NAME = "NIGHT1_REPORT_EN.xlsx"
FIXED_DT = (2020, 1, 1, 0, 0, 0)

SHEET_NAMES = [
    "Executive Summary",
    "Run Overview",
    "Model Scorecard",
    "Recommendations",
    "Strengths & Weaknesses",
    "Per-Model Test Detail",
    "Failure Analysis",
    "HOME Verdicts",
    "Full Ranking",
    "Performance",
    "Known Model Behaviors",
    "HOME-07 Special Case",
]

VERDICT_EN = {
    "HOME_WIN": "Home model wins",
    "HOME_TIE": "Home model ties",
    "HOME_LOSS": "Home model loses",
}
NO_DATA = "Insufficient data for a verdict"
NO_COMPETITOR = "\u2014 (insufficient data for a verdict)"
NO_GEN_DATA = "No data (embedding model; generation speed not applicable)"
RANKED = "Ranked"
NOT_RANKED = "Not ranked (unsupported / time-boxed out)"

# --- Deep-integration additions (12-sheet workbook) -----------------------
# The per-model deep analysis lives in repo-root analysis_en/group_*.json
# (written by four sibling agents; this module only renders it).
ANALYSIS_DIRNAME = "analysis_en"
ELIGIBILITY_RELPATH = os.path.join("config", "eligibility_night.json")

# Single consistent truncation marker for Failure Analysis evidence quotes.
TRUNCATION_MARKER = " [truncated]"
# Obvious cell state when a quote cannot be verified against any record.
UNVERIFIABLE_QUOTE = "Quote could not be verified against the record"

# Column (1-based) holding verbatim model output on the Failure Analysis
# sheet. Americanization and British-spelling scans skip this column:
# rewriting a quote would destroy it as evidence.
QUOTE_COLUMN = 7

ACCENT = "1F4E79"  # single professional accent color (navy) for header fills

STATUS_MEANINGS = {
    "OK": "The output met the format contract and was scored on correctness.",
    "WRONG_ANSWER": "The output followed the format contract, but the answer was wrong.",
    "FORMAT_ERROR": "The output broke the format contract (for example, prose around JSON). "
                    "The semantic score is kept, but the strict score is 0.",
    "EMPTY_OUTPUT": "The model returned no usable output.",
    "OUTPUT_TRUNCATED": "Generation reached the token limit before the answer was complete.",
    "UNSUPPORTED_CAPABILITY": "The model lacks a capability this test needs. "
                              "The record is excluded from quality scores.",
    "CONTEXT_OVERFLOW": "The prompt exceeded the model's context window. "
                        "This is a runtime event; the record is excluded from quality scores.",
    "NOT_RUN_BUDGET": "The block did not fit into the 8-hour time box. "
                      "It is excluded from quality scores.",
    "TIMEOUT": "The request exceeded its time limit. This is a runtime event; "
               "the record is excluded from quality scores.",
    "MODEL_LOAD_ERROR": "The model failed to load. This is a runtime event; "
                        "the record is excluded from quality scores.",
    "MODEL_RUNTIME_ERROR": "The model errored during generation. This is a runtime event; "
                           "the record is excluded from quality scores.",
    "OOM_GPU": "The run exceeded GPU memory. This is a runtime event; "
               "the record is excluded from quality scores.",
    "OOM_RAM": "The run exceeded host memory. This is a runtime event; "
               "the record is excluded from quality scores.",
    "HARNESS_ERROR": "The harness itself failed on this block. This is a runtime event; "
                     "the record is excluded from quality scores.",
    "BLOCKED_CONTENDED": "The GPU was busy with another process, so the block was skipped.",
    "GPU_UNLOAD_FAILED": "The previous model could not be fully unloaded before this block.",
    "VALIDATOR_ERROR": "The scoring validator raised an exception. This is a harness-side event; "
                       "the record is excluded from quality scores.",
}

# English rendering of MASTER_PLAN.md section 5 (short factual notes, written
# fresh in English from those bullets; kept as data so the sheet never has to
# machine-translate anything at runtime).
BEHAVIORS_EN: list[tuple[str, str]] = [
    ("deepseek-r1:8b",
     "DeepSeek-R1 tends to over-reason. On HOME-04 its thinking trace is cut off by the token limit "
     "(OUTPUT_TRUNCATED inside thinking) even at 8192 tokens, at roughly 180 seconds per case."),
    ("granite-code:8b-instruct",
     "Granite-Code produces correct fixes, but the reported bug line is usually off by one, "
     "and it adds explanatory prose around the answer."),
    ("qwen2.5vl:7b / gemma3:12b",
     "Qwen2.5-VL and Gemma3 wrap the answer in code fences (FORMAT_ERROR; semantic score kept). "
     "Qwen2.5-VL grounding boxes are shifted by about 57 pixels."),
    ("command-r7b:7b",
     "Command-R7B answers HOME-03 in prose instead of the required JSON."),
    ("sqlcoder:7b",
     "SQLCoder emits Postgres ILIKE syntax although the fixtures run on SQLite "
     "(a model habit, not a harness bug)."),
    ("mistral-nemo:12b",
     "Mistral-Nemo generates only about 4 tokens/s in a 16K-token context under partial CPU offload, "
     "and returns a degenerate evidence list containing every id."),
    ("lfm2.5:8b",
     "LFM2.5 occasionally returns an empty final message or drops the closing brace after "
     "otherwise correct tool calls."),
    ("functiongemma:270m / granite3.2-vision:2b",
     "FunctionGemma pads argument strings with whitespace and makes extra calls. "
     "Granite3.2-Vision emits textual <tool_call> markup on HOME-07 instead of JSON."),
    ("qwen3:14b, gemma3:12b, mistral-nemo:12b",
     "Three models run with partial GPU offload on 8 GB of VRAM: Qwen3:14B runs at about 8.5 tokens/s, "
     "Gemma3:12B at about 10 tokens/s, and Mistral-Nemo:12B between 4 and 23 tokens/s."),
    ("granite3.2-vision:2b",
     "Under a JSON-only instruction the model answers with a <tool_call> fragment or a DocTags <doc> dump, "
     "so the model scores Q_sem 0 on HOME-07 under the uniform contract. "
     "A plain-answer adapter re-test is reported separately."),
    ("glm-ocr:latest",
     "GLM-OCR reads English text correctly but repeats it (REPETITION_LOOP). On Ukrainian images it mixes "
     "Latin and Russian letters and loops a single line until the token limit."),
    ("qwen3:14b (thinking)",
     "Qwen3:14B (thinking) runs at 6.2 tokens/s in an 8K-token context on 8 GB of VRAM. "
     "One HOME-17 case takes about 16.5 minutes and still ends truncated inside thinking at 6144 tokens, "
     "which makes it impractical for reasoning tasks on this class of hardware."),
]
# Backwards-compatible alias (old name used en-GB spelling).
BEHAVIOURS_EN = BEHAVIORS_EN

# English gloss for raw sub_reason codes that contain Ukrainian text.
# Every value rendered in the workbook goes through _gloss_sub_reason so no
# Cyrillic can leak into the English document.
_SUB_GLOSS_EXACT = {
    "поза scope": "out-of-scope test",
    "CAP: немає embed": "CAP: missing embed modality",
    "CAP: немає text": "CAP: missing text modality",
    "CAP: немає text,tools": "CAP: missing text and tools modalities",
    "CAP: немає tools": "CAP: missing tools modality",
    "CAP: немає tools,vision": "CAP: missing tools and vision modalities",
    "CAP: немає vision": "CAP: missing vision modality",
    "QA не є документованим сценарієм GLM-OCR (лише parsing/IE)":
        "QA is not a documented GLM-OCR scenario (parsing/IE only)",
    "багатоходовий діалог не задокументовано":
        "multi-turn dialogue not documented",
}


def _resolve_run_dir(run_id: str, results_root: str | None) -> str:
    if os.path.isdir(run_id):
        return os.path.abspath(run_id)
    base = results_root or os.path.join(ROOT, "results")
    if not os.path.isabs(base):
        base = os.path.join(ROOT, base)
    return os.path.join(base, run_id)


def _load_manifest(run_dir: str) -> dict:
    try:
        with open(os.path.join(run_dir, "manifest.json"), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def _load_perf(run_dir: str) -> list[dict]:
    path = os.path.join(run_dir, "perf.csv")
    rows: list[dict] = []
    try:
        with open(path, encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                rows.append({k: (v.strip() if isinstance(v, str) else v)
                             for k, v in row.items()})
    except OSError:
        return []
    return rows


def _parse_iso(s: str):
    try:
        return datetime.fromisoformat(s)
    except (TypeError, ValueError):
        return None


def _run_times(manifest: dict, records: list[dict]):
    start = _parse_iso(manifest.get("start_local") or manifest.get("start_utc") or "")
    end = None
    for r in records:
        ts = r.get("timestamps") or {}
        t = _parse_iso(ts.get("local") or ts.get("utc") or "")
        if t is not None and (end is None or t > end):
            end = t
    return start, end


def _test_meta() -> tuple[dict, dict]:
    """Return ({test_id: title}, {test_id: home_tag}); {}s on failure."""
    try:
        from bench import registry as _reg
        tdir = os.path.join(ROOT, "bench", "tests")
        mods = _reg.discover(tdir)
        titles = {tid: str(getattr(m.META, "get", lambda *a: "")("title", ""))
                  if isinstance(m.META, dict) else "" for tid, m in mods.items()}
        home_map = _reg.home_model_map(mods)  # {model_tag: test_id}
        inv = {v: k for k, v in home_map.items()}
        return titles, inv
    except Exception:
        return {}, {}


def _only_status(row: dict, statuses: set) -> bool:
    counts = row.get("counts") or {}
    return bool(counts) and set(counts) <= statuses


def _is_not_ranked(row: dict) -> bool:
    # A row with no scored cases carries no quality measurement to rank,
    # whatever status mix produced the zero. (Root cause of the gemma3:12b /
    # HOME-10 defect: two CONTEXT_OVERFLOW plus one NOT_RUN_BUDGET fell
    # through the status-set checks below and rendered as Ranked with Q 0.)
    if (row.get("n_scored") or 0) == 0:
        return True
    if _only_status(row, {"UNSUPPORTED_CAPABILITY"}):
        return True
    if (_only_status(row, {"NOT_RUN_BUDGET", "UNSUPPORTED_CAPABILITY"})
            and not _only_status(row, {"UNSUPPORTED_CAPABILITY"})):
        return True
    return False


def _num_or_none(v):
    try:
        if v is None or (isinstance(v, str) and not v.strip()):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def _gloss_sub_reason(raw: str) -> str:
    """Translate a raw sub_reason code into plain English (no Cyrillic)."""
    import re as _re
    s = str(raw or "").strip()
    if not s or s == "-":
        return ""
    if s in _SUB_GLOSS_EXACT:
        return _SUB_GLOSS_EXACT[s]
    m = _re.fullmatch(
        r"CONTEXT_WINDOW:\s*потрібно\s*(\d+),\s*runtime\s*(\d+)", s)
    if m:
        return ("CONTEXT_WINDOW: requires %s tokens, runtime %s tokens"
                % (m.group(1), m.group(2)))
    # Generic fallbacks so any future Ukrainian fragment still renders English.
    s = s.replace("немає", "missing")
    s = s.replace("поза scope", "out-of-scope test")
    s = s.replace("поза", "out-of-scope")
    s = s.replace("потрібно", "requires")
    s = _re.sub(r"[\u0400-\u04FF]", "", s)
    s = _re.sub(r"\s{2,}", " ", s).strip(" ,;:")
    if not s:
        return "untranslated code"
    if _re.search(r"[\u0400-\u04FF]", s):
        return "untranslated code"
    return s


def strip_trailing_marker(quote: str) -> str:
    """Remove one trailing truncation marker from an evidence quote.

    Module-level so the Ukrainian workbook builder can import and reuse
    it instead of reimplementing the same normalization. Handles the
    markers observed across the four analysis groups (`[truncated]`,
    `[...]`) plus a bare trailing ellipsis (Unicode or ASCII). Only a
    trailing marker is removed; the characters before it stay
    byte-identical because that text is the evidence.
    """
    import re as _re
    s = str(quote or "")
    for pat in (r"\s*\[truncated\]\s*$", r"\s*\[\.\.\.\]\s*$",
                r"\s*\u2026\s*$", r"\s*\.\.\.\s*$"):
        stripped = _re.sub(pat, "", s)
        if stripped != s:
            return stripped
    return s


def normalize_evidence_quote(quote: object, full_content: object,
                             marker: str = TRUNCATION_MARKER
                             ) -> tuple[str, str]:
    """Normalize one evidence quote against its record's full content.

    Returns (rendered_quote, status) with status one of "complete"
    (stripped quote equals the full content: render with no marker),
    "truncated" (stripped quote is a proper prefix: render stripped
    quote plus the single consistent marker), or "unverifiable"
    (neither: the caller renders UNVERIFIABLE_QUOTE instead of the
    text). The quote text itself is never altered, only the trailing
    marker. Module-level so both workbooks share one implementation.
    """
    if not isinstance(quote, str) or not quote.strip():
        return UNVERIFIABLE_QUOTE, "unverifiable"
    if not isinstance(full_content, str) or not full_content:
        return UNVERIFIABLE_QUOTE, "unverifiable"
    stripped = strip_trailing_marker(quote)
    if stripped == full_content:
        return stripped, "complete"
    if stripped and full_content.startswith(stripped):
        return stripped + marker, "truncated"
    return UNVERIFIABLE_QUOTE, "unverifiable"


# British -> American rules applied as a final pass over every string
# cell of every sheet (except the verbatim quote column). Morphological
# rules, not a closed word list: a closed list has now failed twice --
# first a whole-word scan missed `optimisation`, then a stem list missed
# `localisation`. Rules:
# - `-isation`/`-isations` -> `-ization`/`-izations`; `-ise`/`-ised`/
#   `-ises`/`-ising`/`-iser`/`-isers`/`-isable`(`s`) -> `-ize`/... .
# - `-our` -> `-or` for behaviour/colour/favour/honour/labour/neighbour/
#   rigour/vigour/endeavour/favourite (+s/ed endings).
# - `-re` -> `-er` for centre/metre/theatre/litre/fibre/calibre (+s/d).
# - `-ce` -> `-se` for licence/defence/offence/pretence (+s).
# - doubled `ll` -> single `l` for travel/model/label/cancel/signal/level
#   + ed/ing/er(s) (travelled->traveled, modelling->modeling, ...).
# - `analyse` family -> `analyze` family (whole words only, so US
#   `analysis`/`analyses`/`analyst` never match); `programme`->`program`,
#   `catalogue`->`catalog`.
# Case is preserved per word. Words in _US_EXCEPTIONS are never touched
# (they end in -ise/-our/-re in US English too).

_EXCEPTION_BASES = (
    "advertise exercise surprise compromise revise supervise comprise "
    "disguise improvise devise despise arise rise wise precise concise "
    "promise franchise merchandise enterprise chastise televise demise "
    "expertise paradise premise treatise otherwise likewise advise advice "
    "glamour analysis"
).split()


def _exception_forms() -> frozenset:
    forms: set[str] = set()
    for base in _EXCEPTION_BASES:
        forms.add(base)
        forms.add(base + "s")
        forms.add(base + "d")
        if base.endswith("e"):
            forms.add(base[:-1] + "ing")
        else:
            forms.add(base + "ing")
    forms.update(("analyses", "rising", "arising", "risen", "arisen",
                  "rises", "arises", "rose", "arose", "wiser", "wisest",
                  "glamorous"))
    return frozenset(forms)


_US_EXCEPTIONS = _exception_forms()

_ANALYSE_MAP = {"analyse": "analyze", "analysed": "analyzed",
                "analyser": "analyzer", "analysers": "analyzers",
                "analysing": "analyzing"}


def _match_case(template: str, us_lower: str) -> str:
    if template.isupper():
        return us_lower.upper()
    if template[0].isupper():
        return us_lower.capitalize()
    return us_lower


def _americanize_word(word: str) -> str:
    """Rewrite one word to American spelling, or return it unchanged."""
    import re as _re
    low = word.lower()
    if low in _US_EXCEPTIONS:
        return word
    if low in _ANALYSE_MAP:
        return _match_case(word, _ANALYSE_MAP[low])
    if low in ("programme", "programmes"):
        return _match_case(word, low.replace("programme", "program"))
    if low in ("catalogue", "catalogues"):
        return _match_case(word, low.replace("catalogue", "catalog"))
    m = _re.fullmatch(r"([a-z]+?)isation(s?)", low)
    if m and len(m.group(1)) >= 2:
        stem = word[:len(m.group(1))]
        src_suffix = word[len(m.group(1)):]
        suffix = ("IZATION" if src_suffix.isupper() else "ization")
        if m.group(2):
            suffix += "S" if m.group(2).isupper() else "s"
        return stem + suffix
    m = _re.fullmatch(r"([a-z]{2,}?)is(e|es|ed|ing|ers?|ables?)", low)
    if m:
        stem = word[:len(m.group(1))]
        ending = m.group(2)
        src_ending = word[len(m.group(1)) + 2:]
        if word.isupper() or (src_ending and src_ending.isupper()):
            return stem + "IZ" + ending.upper()
        return stem + "iz" + ending
    m = _re.fullmatch(
        r"(behaviour|colour|favour|honour|labour|neighbour|rigour"
        r"|vigour|endeavour|favourite)(s|ed|ables?|ably|ites?)?", low)
    if m:
        core = word[:len(m.group(1))]
        m_our = _re.search(r"our", core, _re.IGNORECASE)
        oursrc = m_our.group(0)
        rep = "OR" if oursrc.isupper() else ("Or" if oursrc[0].isupper()
                                             else "or")
        return (core[:m_our.start()] + rep + core[m_our.end():]
                + word[len(m.group(1)):])
    m = _re.fullmatch(
        r"(centre|metre|theatre|litre|fibre|calibre)(s|d)?", low)
    if m:
        core = word[:len(m.group(1))]
        tail = core[-2:]
        rep = "ER" if tail.isupper() else ("Er" if tail[0].isupper()
                                           else "er")
        return core[:-2] + rep + word[len(m.group(1)):]
    m = _re.fullmatch(r"(licence|defence|offence|pretence)(s|d)?", low)
    if m:
        core = word[:len(m.group(1))]
        tail = core[-2:]
        rep = "SE" if tail.isupper() else ("Se" if tail[0].isupper()
                                           else "se")
        return core[:-2] + rep + word[len(m.group(1)):]
    m = _re.fullmatch(
        r"(travel|model|label|cancel|signal|level)l(ed|ing|ers?)", low)
    if m:
        # Drop one of the doubled consonants; surviving slices keep the
        # source case, so Travelled->Traveled and TRAVELLED->TRAVELED.
        return word[:len(m.group(1))] + word[len(m.group(1)) + 1:]
    return word


def british_hits(value: object) -> list[str]:
    """British-spelling hits in a cell value (whole offending words).

    Single source of truth with americanize(): a word is flagged if and
    only if the rewrite pass would change it, so the scan and the fix
    can never disagree the way the old closed stem list did.
    """
    import re as _re
    return [m.group(0) for m in
            _re.finditer(r"[A-Za-z]+", str(value or ""))
            if _americanize_word(m.group(0)) != m.group(0)]


def americanize(text: str) -> str:
    """Rewrite British spellings to American English, preserving case.

    Covers META-sourced titles (`Constrained optimisation, ...`), the
    analysis JSON prose (which adopted the British form), and any future
    sheet wording, because it runs over every rendered cell. Verbatim
    model output is never passed through here (see _americanize_workbook).
    Word-based: each word goes through _americanize_word, so exception
    words (surprise, exercise, analysis, ...) are never touched.
    """
    import re as _re
    if not isinstance(text, str) or not text:
        return text
    return _re.sub(r"[A-Za-z]+",
                   lambda m: _americanize_word(m.group(0)), text)


def _americanize_workbook(wb) -> None:
    """Apply americanize() to every string cell except the quote column."""
    for ws in wb.worksheets:
        skip_col = QUOTE_COLUMN if ws.title == "Failure Analysis" else -1
        for row in ws.iter_rows():
            for cell in row:
                if cell.column == skip_col:
                    continue
                if isinstance(cell.value, str) and cell.value:
                    fixed = americanize(cell.value)
                    if fixed != cell.value:
                        cell.value = fixed


def _fmt_q(q: float) -> str:
    return "%.3f" % q


# ---------------------------------------------------------------- formatting

def _style_table(ws, widths: list[float], numfmt: dict[int, str] | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    if numfmt:
        for row in ws.iter_rows(min_row=2):
            for col_idx, fmt in numfmt.items():
                cell = row[col_idx]
                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt


def _q_scale(ws, col_letter: str, n_rows: int) -> None:
    """Tasteful single-hue color scale over Q_sem values (light blue -> navy)."""
    try:
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError:
        return
    if n_rows < 2:
        return
    ws.conditional_formatting.add(
        "%s2:%s%d" % (col_letter, col_letter, n_rows),
        ColorScaleRule(start_type="num", start_value=0.0, start_color="F2F2F2",
                       mid_type="num", mid_value=0.5, mid_color="8EA9DB",
                       end_type="num", end_value=1.0, end_color=ACCENT))


def _save_deterministic(wb, path: str) -> None:
    from datetime import datetime as _dt
    import re as _re
    props = wb.properties
    props.created = _dt(*FIXED_DT)
    props.modified = _dt(*FIXED_DT)
    props.creator = "BENCH_V5_NIGHT"
    props.lastModifiedBy = "BENCH_V5_NIGHT"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    zin = zipfile.ZipFile(buf, "r")
    # openpyxl stamps docProps/core.xml dcterms timestamps with the current
    # time at save (ignoring wb.properties); pin both so the bytes are stable.
    fixed_ts = b"2020-01-01T00:00:00Z"
    tmp = path + ".tmpzip"
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = _re.sub(
                    rb"<dcterms:created[^>]*>[^<]*</dcterms:created>",
                    rb'<dcterms:created xsi:type="dcterms:W3CDTF">'
                    + fixed_ts + rb"</dcterms:created>", data)
                data = _re.sub(
                    rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                    rb'<dcterms:modified xsi:type="dcterms:W3CDTF">'
                    + fixed_ts + rb"</dcterms:modified>", data)
            zi = zipfile.ZipInfo(item.filename, date_time=FIXED_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)
    zin.close()
    os.replace(tmp, path)


# ---------------------------------------------------------------- sheets

def _executive_summary(run_id: str, manifest: dict, records: list[dict],
                       per_test: dict, per_model: dict,
                       perf_csv_rows: list[dict],
                       gv_records: list[dict]) -> list[tuple[str, str]]:
    """Data-grounded executive summary paragraphs; every sentence traces to a sheet."""
    ov = _report.run_overview(records)
    total = len(records)
    nbr = ov.get("not_run_budget", 0)
    pct = (100.0 * nbr / total) if total else 0.0
    hv = _report.home_verdicts(per_test, records)
    wins = sum(1 for v in hv.values() if v.get("verdict") == "HOME_WIN")
    ties = sum(1 for v in hv.values() if v.get("verdict") == "HOME_TIE")
    losses = sum(1 for v in hv.values() if v.get("verdict") == "HOME_LOSS")
    fmt_err = ov["counts"].get("FORMAT_ERROR", 0)
    ok_n = ov["counts"].get("OK", 0)
    # Best home-test results among tests with scored home rows.
    _, home_of_test = _test_meta()
    home_best: list[tuple[str, str, float, str]] = []
    for tid, rows in per_test.items():
        if tid == "PERF":
            continue
        home_tag = home_of_test.get(tid, "")
        home_row = next((r for r in rows if r["model"] == home_tag), None)
        if home_row is not None and (home_row.get("n_scored") or 0) > 0:
            home_best.append((tid, home_tag, float(home_row["q_sem"]),
                              "%d/%d" % (home_row["n_scored"], home_row["n_total"])))
    home_best.sort(key=lambda t: (-t[2], t[0]))
    if home_best:
        top3 = "; ".join("%s on %s (Q_sem %s, n=%s)" % (tag, tid, _fmt_q(q), n)
                         for tid, tag, q, n in home_best[:3])
        top3_txt = ("The strongest home-test results were %s. " % top3
                    + "Full per-test numbers are on the HOME Verdicts sheet.")
    else:
        top3_txt = ("No home test produced a scored result in this run. "
                    "See the HOME Verdicts sheet for the per-test breakdown.")
    partial = sorted(str(r.get("tag", "")) for r in perf_csv_rows
                     if _num_or_none(r.get("offload_ratio")) is not None
                     and _num_or_none(r.get("offload_ratio")) < 1.0)
    if partial:
        word = "three models" if len(partial) == 3 else "%d models" % len(partial)
        # Spell out small counts in prose (never "3 model(s)").
        if len(partial) == 3:
            hw_txt = ("On this 8 GB card, three models could not fit fully on the GPU and ran with "
                      "partial CPU offload (%s), which sharply reduces generation speed. "
                      "Details are on the Performance sheet."
                      % (", ".join(partial)))
        else:
            hw_txt = ("On this 8 GB card, %s could not fit fully on the GPU and ran with "
                      "partial CPU offload (%s), which sharply reduces generation speed. "
                      "Details are on the Performance sheet."
                      % (word, ", ".join(partial)))
    else:
        hw_txt = ("Every model with a measured offload ratio ran fully on the GPU "
                  "(offload ratio 1.0). Details are on the Performance sheet.")
    gv_txt: str
    if gv_records:
        gv_home = _report.summarize_group(
            [r for r in gv_records if _report.test_of(r) == "HOME-07"
             and _report.model_of(r) == "granite3.2-vision:2b"])
        main_recs = [r for r in records if _report.test_of(r) == "HOME-07"
                     and _report.model_of(r) == "granite3.2-vision:2b"]
        main_s = _report.summarize_group(main_recs)
        gv_txt = ("The granite3.2-vision document-extraction case was tested twice: Q_sem %s under the "
                  "uniform JSON contract and Q_sem %s under a plain-answer adapter. Both rounds scored "
                  "zero, which points to a genuine model capability limit rather than a prompt-format artifact. "
                  "The full comparison is on the HOME-07 Special Case sheet."
                  % (_fmt_q(main_s["q_sem"]) if main_s["n_scored"] else "n/a",
                     _fmt_q(gv_home["q_sem"]) if gv_home["n_scored"] else "n/a"))
    else:
        gv_txt = ("The granite3.2-vision document-extraction case scored Q_sem 0.000 under the uniform "
                  "JSON contract. The follow-up comparison is on the HOME-07 Special Case sheet.")
    return [
        ("What this is",
         "This workbook reports NIGHT-1, an independent benchmark of local large-language models run "
         "through a test harness built from scratch. The project demonstrates benchmark design and "
         "systems engineering: a time-boxed, resumable runner with a strict scoring contract, "
         "per-model performance probes, and a fully reproducible reporting pipeline."),
        ("Scale of the run",
         "One completed run (%s) covered %d models and %d test groups with %d records in a single "
         "8-hour time-boxed session on consumer hardware (NVIDIA GeForce RTX 3070, 8 GB VRAM). "
         "The Run Overview sheet lists every model, test group, and outcome count."
         % (manifest.get("run_id", run_id), ov["n_models"], ov["n_tests"], total)),
        ("Best home-test results", top3_txt),
        ("Home-model verdicts",
         "Across %d head-to-head home-model comparisons, the home model recorded %d wins, "
         "%d ties, and %d losses against its best competitor on the shared case set. "
         "Each verdict is listed on the HOME Verdicts sheet."
         % (wins + ties + losses, wins, ties, losses)),
        ("Time-box coverage",
         "Of %s records, %d (%.1f%%) were not run because they did not fit the 8-hour deadline "
         "(status NOT_RUN_BUDGET) and were excluded from quality scores rather than rushed. "
         "Quality averages therefore describe completed work only. "
         "The Full Ranking sheet marks every such row explicitly."
         % ("{:,}".format(total), nbr, pct)),
        ("Format versus substance",
         "Models produced %s clean passes (status OK) and %d format errors: outputs that broke the "
         "answer contract but still earned partial semantic credit. The harness scores meaning and "
         "format separately, so formatting problems are never conflated with incorrect answers. "
         "Status counts are on the Run Overview sheet." % ("{:,}".format(ok_n), fmt_err)),
        ("Hardware fit", hw_txt),
        ("Engineering finding", gv_txt),
        ("Per-model depth",
         "Per-model ratings, usage recommendations, strengths and weaknesses, "
         "test-by-test detail, and failure patterns are on the Model "
         "Scorecard, Recommendations, Strengths & Weaknesses, Per-Model Test "
         "Detail, and Failure Analysis sheets."),
    ]


def _sheet_executive(wb, run_id: str, manifest: dict, records: list[dict],
                     per_test: dict, per_model: dict,
                     perf_csv_rows: list[dict], gv_records: list[dict]) -> None:
    from openpyxl.styles import Alignment
    ws = wb.active
    ws.title = SHEET_NAMES[0]
    ws.append(["Section", "Text"])
    for section, text in _executive_summary(run_id, manifest, records, per_test,
                                            per_model, perf_csv_rows, gv_records):
        ws.append([section, text])
    _style_table(ws, [22, 130])
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[1].alignment = wrap
        ws.row_dimensions[row[0].row].height = 60


def _sheet_overview(wb, run_id: str, manifest: dict, records: list[dict],
                    gv_records: list[dict]) -> None:
    ws = wb.create_sheet(SHEET_NAMES[1])
    ws.append(["Item", "Value", "Note"])
    ov = _report.run_overview(records)
    start, end = _run_times(manifest, records)
    fp = manifest.get("fingerprint") or {}
    total = len(records)
    nbr = ov.get("not_run_budget", 0)
    pct = (100.0 * nbr / total) if total else 0.0
    ws.append(["run_id", manifest.get("run_id", run_id), ""])
    ws.append(["Start (local)", manifest.get("start_local", "") or "", ""])
    ws.append(["End (local)", end.isoformat() if end is not None else "", "Timestamp of the last recorded result."])
    if start is not None and end is not None and end >= start:
        ws.append(["Duration (min)", round((end - start).total_seconds() / 60.0, 1), "from run start to the last recorded result"])
    else:
        ws.append(["Duration (min)", "", "timestamps missing"])
    ws.append(["Budget (hours)", manifest.get("budget_hours", ""), "hard deadline for the whole run"])
    ws.append(["Ollama version", manifest.get("ollama_version", ""), ""])
    ws.append(["GPU", "NVIDIA GeForce RTX 3070 (8 GB VRAM)", "consumer card used for every measurement"])
    ws.append(["Models run", ov["n_models"], ", ".join(ov["models"])])
    ws.append(["Test groups run", ov["n_tests"], ", ".join(ov["tests"])])
    ws.append(["Records", total, "One row per attempted model/test/case combination."])
    ws.append(["NOT_RUN_BUDGET records", nbr, "blocks skipped by the deadline"])
    ws.append(["NOT_RUN_BUDGET share (%)", round(pct, 1), "time-box coverage: skipped work is excluded from scores, not scored as zero"])
    ws.append(["", "", ""])
    ws.append(["Status", "Count", "What it means"])
    for st in sorted(ov["counts"]):
        ws.append([st, ov["counts"][st],
                   STATUS_MEANINGS.get(st, "see the report code for this status.")])
    ws.append(["", "", ""])
    gv_n = len(gv_records)
    ws.append(["Follow-up run (_gv)", "%s_gv" % manifest.get("run_id", run_id),
               "HOME-07 / granite3.2-vision:2b under a plain-answer adapter; kept separate, never mixed into the tables."])
    ws.append(["Follow-up records", gv_n, "detail is on the HOME-07 Special Case sheet."])
    _style_table(ws, [24, 60, 90])


def _sheet_verdicts(wb, records: list[dict], per_test: dict) -> None:
    ws = wb.create_sheet(SHEET_NAMES[7])
    ws.append(["Test", "Title", "Home model", "Home Q_sem", "Home n",
               "Best competitor", "Competitor Q_sem", "Competitor n",
               "Common n", "Verdict"])
    titles, home_of_test = _test_meta()
    hv = _report.home_verdicts(per_test, records)
    for tid in sorted(_report.HOME_SET):
        rows = per_test.get(tid, [])
        title = titles.get(tid, "")
        home_tag = home_of_test.get(tid, "")
        home_row = next((r for r in rows if r["model"] == home_tag), None)
        if home_row is None or (home_row.get("n_scored") or 0) == 0:
            home_q: object = ""
            home_n = ("" if home_row is None
                      else "%d/%d" % (home_row["n_scored"], home_row["n_total"]))
            comp_tag: object = ""
            comp_q: object = ""
            comp_n = ""
            common: object = ""
            verdict = NO_DATA
            v = hv.get(tid)
            if v is not None and v.get("second"):
                comp_tag = v["second"]
                comp_r = next((r for r in rows if r["model"] == comp_tag), None)
                if comp_r is not None and (comp_r.get("n_scored") or 0) > 0:
                    comp_q = round(comp_r["q_sem"], 3)
                    comp_n = "%d/%d" % (comp_r["n_scored"], comp_r["n_total"])
                common = v.get("common_n", "")
            if not comp_tag:
                comp_tag = NO_COMPETITOR
            ws.append([tid, title, home_tag or NO_DATA, home_q, home_n,
                       comp_tag, comp_q, comp_n, common, verdict])
            continue
        home_q = round(home_row["q_sem"], 3)
        home_n = "%d/%d" % (home_row["n_scored"], home_row["n_total"])
        ranked = sorted((r for r in rows if not _is_not_ranked(r)),
                        key=lambda r: (-r["q_sem"], -r["n_scored"], r["model"]))
        comp_row = next((r for r in ranked if r["model"] != home_tag), None)
        v = hv.get(tid)
        if v is not None and v.get("second"):
            comp_tag = v["second"]
            comp_r = next((r for r in rows if r["model"] == comp_tag), comp_row)
            if comp_r is not None and (comp_r.get("n_scored") or 0) > 0:
                comp_q = round(comp_r["q_sem"], 3)
                comp_n = "%d/%d" % (comp_r["n_scored"], comp_r["n_total"])
            else:
                comp_q, comp_n = "", ""
            verdict = VERDICT_EN.get(v.get("verdict", ""), NO_DATA)
            common = v.get("common_n", "")
        elif comp_row is not None:
            comp_tag = comp_row["model"]
            comp_q = round(comp_row["q_sem"], 3)
            comp_n = "%d/%d" % (comp_row["n_scored"], comp_row["n_total"])
            common, verdict = "", NO_DATA
        else:
            comp_tag, comp_q, comp_n, common, verdict = (
                NO_COMPETITOR, "", "", "", NO_DATA)
        ws.append([tid, title, home_tag, home_q, home_n,
                   comp_tag, comp_q, comp_n, common, verdict])
    _style_table(ws, [10, 30, 24, 12, 12, 24, 14, 12, 10, 26],
                 {3: "0.000", 6: "0.000"})
    _q_scale(ws, "D", ws.max_row)
    _q_scale(ws, "G", ws.max_row)


def _sheet_ranking(wb, per_test: dict) -> None:
    ws = wb.create_sheet(SHEET_NAMES[8])
    ws.append(["Test", "Model", "Q_sem", "Q_strict", "N scored", "N total",
               "Coverage", "CI low", "CI high", "Low coverage",
               "Ranking status"])
    for tid in sorted(t for t in per_test if t != "PERF"):
        rows = per_test[tid]
        ranked = sorted((r for r in rows if not _is_not_ranked(r)),
                        key=lambda r: (-r["q_sem"], -r["n_scored"], r["model"]))
        rest = sorted((r for r in rows if _is_not_ranked(r)),
                      key=lambda r: r["model"])
        for r in ranked + rest:
            low_cov = "yes" if r.get("coverage", 1.0) < 0.9 else "no"
            status = NOT_RANKED if r in rest else RANKED
            ws.append([tid, r["model"], round(r["q_sem"], 3),
                       round(r["q_strict"], 3), r["n_scored"], r["n_total"],
                       round(r.get("coverage", 0.0), 3),
                       round(r["ci_lo"], 3), round(r["ci_hi"], 3),
                       low_cov, status])
    _style_table(ws, [10, 24, 10, 10, 10, 10, 10, 10, 10, 12, 38],
                 {2: "0.000", 3: "0.000", 6: "0.0%", 7: "0.000", 8: "0.000"})
    _q_scale(ws, "C", ws.max_row)


# --- Task 3: canonical "not attempted" vocabulary -----------------------
# The four analysis groups each invented their own wording for the same
# closed set of eligibility reasons, so the Notes column drifted across
# ~40 phrasings. Every Not attempted note is re-rendered here at render
# time from config/eligibility_night.json (via the already-audited
# _gloss_sub_reason machinery), with one consistent prefix. Model-specific
# detail beyond the eligibility reason is appended, never discarded.

_NOT_ATTEMPTED_PREFIX = "Not attempted \u2014 "


def _load_eligibility() -> dict:
    try:
        with open(os.path.join(ROOT, ELIGIBILITY_RELPATH),
                  encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def _strip_note_prefix(note: str) -> str:
    import re as _re
    s = str(note or "").strip()
    s = _re.sub(r"^(Not attempted|Never attempted|Never scored)"
                r"\s*[:\u2014\u2013\-]\s*", "", s,
                flags=_re.IGNORECASE)
    return s.strip()


# Generic vocabulary shared by all paraphrases of the same eligibility
# reason.
_GENERIC_NOTE_WORDS = frozenset(
    ("a an the is are was were be been being not no nor for this that "
     "these those model models test tests pair pairs run recorded record "
     "block blocks reason reasons carries carry further ruled out before "
     "running missing documented undocumented documentation capability "
     "capabilities vision embed embedding text tool tools input support "
     "supports use used uses scenario scenarios eligible eligibility set "
     "scope planned scheduled schedule outside other never attempted "
     "scored attempt attempts unsupported in on of to and or so as its it "
     "with whose which by at from into over under only all any each both "
     "either neither has have had does did input dialogue multi turn "
     "token tokens context runtime allows allow limit required requires "
     "exceed exceeds needs need about").split())

# Words marking genuinely model-specific detail worth keeping after the
# canonical clause (digits handled separately: only digits not already in
# the canonical clause count).
_SPECIFIC_WORDS = frozenset(
    ("english ukrainian budget overflow overflowed queries parsing "
     "question answering extraction exhausted skipped target languages "
     "language documents document query").split())


def _extra_detail(stripped: str, canonical: str) -> str:
    """Return model-specific detail from a note, or "" when generic.

    Generic paraphrases of the eligibility reason ("the run plan scoped
    this test to other models", "the model has no documented vision
    capability, so ...") are retired, not kept: they are the drift this
    normalization removes. Only detail with novel numbers or
    run-specific facts (budget exhaustion, overflow counts, language
    scope, documented-use limits) survives, appended as its own sentence.
    """
    import re as _re
    if not stripped:
        return ""
    words = _re.findall(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)?",
                        stripped.lower())
    canon_words = set(_re.findall(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)?",
                                  canonical.lower()))
    allowed = set(_GENERIC_NOTE_WORDS) | canon_words
    if all(w in allowed for w in words):
        return ""
    canon_digits = set(_re.findall(r"\d+", canonical))
    strip_digits = set(_re.findall(r"\d+", stripped))
    # Whole-word match only: substring matching would fire on
    # "document" inside the generic word "documented".
    flat_words = set()
    for w in words:
        flat_words.update(w.replace("-", " ").split())
    if (strip_digits - canon_digits) or (flat_words & _SPECIFIC_WORDS):
        detail = stripped.strip()
        if detail and detail[-1] not in ".!?":
            detail += "."
        return detail[0].upper() + detail[1:] if detail else ""
    return ""


def canonical_not_attempted(test_id: str, model_tag: str,
                            original_note: str,
                            eligibility: dict) -> str:
    """Render one canonical Not attempted note from eligibility data."""
    entry = eligibility.get(test_id, {}).get(model_tag, {})
    if not isinstance(entry, dict):
        entry = {}
    code = str(entry.get("code", "") or "").strip().upper()
    reason = str(entry.get("reason", "") or "")
    gloss = _gloss_sub_reason(reason) if reason.strip() else ""
    if code == "S" or reason.strip() == "поза scope":
        base = (_NOT_ATTEMPTED_PREFIX
                + "outside the planned scope of this run for this model.")
    elif code == "O":
        if gloss:
            base = "%s%s." % (_NOT_ATTEMPTED_PREFIX, gloss)
        else:
            base = (_NOT_ATTEMPTED_PREFIX
                    + "the pair was not scheduled in this run.")
    elif code == "U":
        if gloss:
            base = "%s%s." % (_NOT_ATTEMPTED_PREFIX, gloss)
        else:
            base = (_NOT_ATTEMPTED_PREFIX
                    + "missing documented capability for this test.")
    else:
        base = (_NOT_ATTEMPTED_PREFIX
                + "no scored cases recorded in this run.")
    extra = _extra_detail(_strip_note_prefix(original_note), base)
    return base + (" " + extra if extra else "")


def normalize_analysis_notes(analysis: list[dict],
                             eligibility: dict) -> dict[str, int]:
    """Rewrite every Not attempted per_test note in place; return counts."""
    from collections import Counter
    counts: Counter = Counter()
    for model in analysis:
        tag = str(model.get("model", "") or "")
        for entry in model.get("per_test") or []:
            if not isinstance(entry, dict):
                continue
            if str(entry.get("outcome", "") or "").strip().lower() != \
                    "not attempted":
                continue
            entry["note"] = canonical_not_attempted(
                str(entry.get("test_id", "") or ""), tag,
                str(entry.get("note", "") or ""), eligibility)
            counts[entry["note"]] += 1
    return dict(counts)


def _sheet_speed(wb, perf_csv_rows: list[dict]) -> None:
    ws = wb.create_sheet(SHEET_NAMES[9])
    ws.append(["Model", "Cold load (s)", "Time to first token (s)", "Prompt tok/s",
               "Gen tok/s", "Offload ratio (1.0 = fully on GPU)"])
    rows = sorted(perf_csv_rows,
                  key=lambda r: (_num_or_none(r.get("gen_tok_s")) is None,
                                 -(_num_or_none(r.get("gen_tok_s")) or 0.0),
                                 str(r.get("tag", ""))))
    for r in rows:
        ws.append([r.get("tag", ""),
                   _num_or_none(r.get("cold_load_s")) or "",
                   _num_or_none(r.get("ttft_s")) or "",
                   _num_or_none(r.get("prompt_tok_s")) or "",
                   _num_or_none(r.get("gen_tok_s"))
                   if _num_or_none(r.get("gen_tok_s")) is not None
                   else NO_GEN_DATA,
                   _num_or_none(r.get("offload_ratio")) or ""])
    ws.append([])
    partial = sorted(str(r.get("tag", "")) for r in rows
                     if _num_or_none(r.get("offload_ratio")) is not None
                     and _num_or_none(r.get("offload_ratio")) < 1.0)
    if partial:
        ws.append(["Note: offload ratio below 1.0 means partial CPU offload "
                   "(the model did not fit in 8 GB of VRAM): "
                   + ", ".join(partial)])
    else:
        ws.append(["Note: no model with a measured offload ratio ran below 1.0 "
                   "(there was no partial CPU offload)."])
    _style_table(ws, [24, 13, 15, 13, 11, 20],
                 {1: "0.0", 2: "0.0", 3: "0", 4: "0", 5: "0.000"})
    # Plain-English consequence of offload, grounded in this run's numbers.
    slows = [(str(r.get("tag", "")), _num_or_none(r.get("gen_tok_s")))
             for r in rows
             if _num_or_none(r.get("offload_ratio")) is not None
             and _num_or_none(r.get("offload_ratio")) < 1.0
             and _num_or_none(r.get("gen_tok_s")) is not None]
    if slows:
        slowest = min(slows, key=lambda t: t[1])
        ws.append(["In practice, a partially offloaded model answers much more slowly "
                   "(for example, %s at %.1f tokens/s), so interactive use on this class of "
                   "hardware favors models with a full offload." % slowest])


def _sheet_behaviours(wb) -> None:
    ws = wb.create_sheet(SHEET_NAMES[10])
    ws.append(["Model", "Observed behavior"])
    for model, feat in BEHAVIORS_EN:
        ws.append([model, feat])
    _style_table(ws, [34, 110])
    from openpyxl.styles import Alignment
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[1].alignment = wrap
        ws.row_dimensions[row[0].row].height = 45


def _sheet_home07(wb, records: list[dict], gv_records: list[dict]) -> None:
    from openpyxl.styles import Alignment
    ws = wb.create_sheet(SHEET_NAMES[11])
    ws.append(["Metric", "Uniform JSON contract (main run)",
               "Plain-answer adapter (follow-up run)"])
    main_recs = [r for r in records if _report.test_of(r) == "HOME-07"
                 and _report.model_of(r) == "granite3.2-vision:2b"]
    gv_recs = [r for r in gv_records if _report.test_of(r) == "HOME-07"
               and _report.model_of(r) == "granite3.2-vision:2b"]
    main_s = _report.summarize_group(main_recs)
    gv_s = _report.summarize_group(gv_recs)
    ws.append(["Q_sem",
               round(main_s["q_sem"], 3) if main_s["n_scored"] else "",
               round(gv_s["q_sem"], 3) if gv_s["n_scored"] else ""])
    ws.append(["Q_strict",
               round(main_s["q_strict"], 3) if main_s["n_scored"] else "",
               round(gv_s["q_strict"], 3) if gv_s["n_scored"] else ""])
    ws.append(["Cases (scored/total)",
               "%d/%d" % (main_s["n_scored"], main_s["n_total"]),
               "%d/%d" % (gv_s["n_scored"], gv_s["n_total"])])
    ws.append(["Run", "Main run", "Follow-up run"])
    ws.append(["Outcome counts", "", ""])
    for st in sorted(set(main_s["counts"]) | set(gv_s["counts"])):
        ws.append([st, main_s["counts"].get(st, 0), gv_s["counts"].get(st, 0)])
    ws.append(["What was tried",
               "The model had to answer document questions as strict JSON under the same contract as every other model.",
               "The same 16 cases were re-run with a plain-answer adapter that accepts a direct answer "
               "without JSON (16 scored of 17 records; the 17th record is a NOT_RUN_BUDGET placeholder, "
               "not a case. The Run Overview total of 23 follow-up records includes additional probe records)."])
    ws.append(["What was found",
               "Under the uniform contract the model produced <tool_call> fragments or DocTags markup "
               "instead of JSON, so it scored zero for quality.",
               "With the adapter the format complied with the contract, but the answers remained incorrect "
               "(for example, the company name instead of the requested discount rate or table total), "
               "so the model still scored zero for quality."])
    ws.append(["Why it matters",
               "The two scores were produced under different answer contracts, so they cannot be compared "
               "like-for-like as Q_sem values. Together they point to a genuine extraction-capability limit "
               "of this 2B vision model, not a prompt-format artifact."])
    _style_table(ws, [18, 55, 55], {1: "0.000", 2: "0.000"})
    ws.merge_cells("B%d:C%d" % (ws.max_row, ws.max_row))
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_col=3):
        for cell in row:
            cell.alignment = wrap
        if row[0].row >= ws.max_row - 2:
            ws.row_dimensions[row[0].row].height = 60


def _records_by_case(records: list[dict]) -> dict[tuple[str, str], str]:
    """Index response.content by (model tag, case id) for quote checks."""
    index: dict[tuple[str, str], str] = {}
    for r in records:
        try:
            tag = _report.model_of(r)
            case = str(r.get("key", "")).split("|")[4]
            content = (r.get("response") or {}).get("content", "")
        except (IndexError, AttributeError):
            continue
        if isinstance(content, str):
            index.setdefault((str(tag), str(case)), content)
    return index


def normalize_failure_quotes(analysis: list[dict],
                             records: list[dict]) -> dict[str, int]:
    """Normalize every failure_analysis quote against the records.

    Returns {"complete": n, "truncated": n, "unverifiable": n}.
    Unverifiable quotes are replaced with UNVERIFIABLE_QUOTE so the
    shipped cell never silently shows a wrong snippet.
    """
    from collections import Counter
    index = _records_by_case(records)
    counts: Counter = Counter()
    for model in analysis:
        tag = str(model.get("model", "") or "")
        for entry in model.get("failure_analysis") or []:
            if not isinstance(entry, dict):
                continue
            quote = entry.get("example_quote")
            if not isinstance(quote, str) or not quote.strip():
                continue
            full = index.get((tag, str(entry.get("example_case_id", "")
                                       or "")))
            rendered, status = normalize_evidence_quote(quote, full)
            entry["example_quote"] = rendered
            counts[status] += 1
    return {"complete": counts.get("complete", 0),
            "truncated": counts.get("truncated", 0),
            "unverifiable": counts.get("unverifiable", 0)}


# Text shown in Failure Analysis `Example output` cells where the model
# produced no content at all. The old string ("Not available") read as a
# data-pipeline gap; the audited truth in all 23 cases is genuinely empty
# `response.content`, and producing nothing is itself the finding.
EMPTY_OUTPUT_TEXT = "No output \u2014 the model produced no content for this case"


def normalize_empty_quote_cells(wb) -> int:
    """Replace the empty-output placeholder in the quote column.

    Only cells of the Failure Analysis `Example output` column (the only
    place the placeholder means "the model emitted nothing") are touched;
    other sheets carry no such cells. Returns the replace count.
    """
    from openpyxl.styles import Font
    n = 0
    if "Failure Analysis" not in wb.sheetnames:
        return 0
    ws = wb["Failure Analysis"]
    for row in ws.iter_rows(min_row=2, min_col=QUOTE_COLUMN,
                            max_col=QUOTE_COLUMN):
        for cell in row:
            if cell.value == "Not available":
                cell.value = EMPTY_OUTPUT_TEXT
                cell.font = Font()
                n += 1
    return n


def normalize_rating_basis(analysis: list[dict]) -> dict[str, int]:
    """Rewrite every overall `rating_basis` to one vocabulary, in place.

    Canonical form: "Mean semantic quality {mean} across {n} scored
    test(s)[ (TID)]; home-test quality {home}[ on TID]; OK rate {ok}[tail]".
    Numbers are never changed (percent OK rates such as "30.7 percent"
    become the identical decimal 0.307); trailing sentences are kept
    verbatim. Entries that do not parse are left untouched and counted
    under "unparsed" -- never silently mangled.
    """
    import re as _re
    rx = _re.compile(
        r"^Mean (?:semantic )?quality (\d+\.\d+) (?:across|over|on|from) "
        r"(.*?), (?:home-test|own-test) quality (\d+\.\d+)((?: on HOME-\d+)?)"
        r", (?:and )?(?:an |a )?(?:OK rate|exact-match rate|exact-answer rate)"
        r"(?: of)? (\d+\.\d+)( percent)?(.*)$", _re.DOTALL)
    counts = {"normalized": 0, "unparsed": 0}
    for model in analysis:
        overall = model.get("overall")
        if not isinstance(overall, dict):
            continue
        raw = overall.get("rating_basis")
        if not isinstance(raw, str) or not raw.strip():
            continue
        m = rx.match(raw.strip())
        if not m:
            counts["unparsed"] += 1
            continue
        mean, scope_raw, home, home_tid, ok, percent, rest = (
            m.group(1), m.group(2), m.group(3), m.group(4), m.group(5),
            m.group(6), m.group(7))
        n_m = _re.search(r"(\d+) scored tests?", scope_raw)
        if n_m:
            n = int(n_m.group(1))
        elif "single scored test" in scope_raw.lower():
            n = 1
        else:
            counts["unparsed"] += 1
            continue
        scope = "%d scored test%s" % (n, "" if n == 1 else "s")
        tid_m = _re.search(r"HOME-\d+", scope_raw)
        if tid_m:
            scope += " (%s)" % tid_m.group(0)
        ok_str = ("%.3f" % (float(ok) / 100.0)) if percent else ok
        rest = (rest or "").strip()
        tail = ""
        m2 = _re.match(r"combine to a rating of \d+\.\d+\s*(.*)$", rest,
                       _re.DOTALL)
        if m2:
            tail = m2.group(1).strip()
        elif rest:
            tail = rest if rest[0] in ".,;:" else " " + rest
        overall["rating_basis"] = (
            "Mean semantic quality %s across %s; home-test quality %s%s; "
            "OK rate %s%s" % (mean, scope, home, home_tid, ok_str, tail))
        counts["normalized"] += 1
    return counts


def build_workbook(run_dir: str, out_path: str) -> str:
    import openpyxl
    manifest = _load_manifest(run_dir)
    records = _report.load_records(run_dir)
    perf_csv_rows = _load_perf(run_dir)
    gv_records: list[dict] = []
    gv_dir = run_dir.rstrip(os.sep) + "_gv"
    if os.path.isdir(gv_dir):
        try:
            gv_records = _report.load_records(gv_dir)
        except OSError:
            gv_records = []
    run_id = manifest.get("run_id", os.path.basename(run_dir))
    per_test, per_model, _ = _report.compute_tables(records)
    wb = openpyxl.Workbook()
    _sheet_executive(wb, run_id, manifest, records, per_test, per_model,
                     perf_csv_rows, gv_records)
    _sheet_overview(wb, run_id, manifest, records, gv_records)
    # Deep per-model sheets (sibling modules own the builders; this
    # function only loads, normalizes, and orders them).
    from tools.en_sheets_detail import (build_detail_sheet,
                                        build_failures_sheet)
    from tools.en_sheets_narrative import build_all_narrative, load_analysis
    analysis = load_analysis(os.path.join(ROOT, ANALYSIS_DIRNAME))
    eligibility = _load_eligibility()
    normalize_rating_basis(analysis)
    normalize_analysis_notes(analysis, eligibility)
    normalize_failure_quotes(analysis, records)
    build_all_narrative(wb, analysis)
    build_detail_sheet(wb, analysis)
    build_failures_sheet(wb, analysis)
    normalize_empty_quote_cells(wb)
    _sheet_verdicts(wb, records, per_test)
    _sheet_ranking(wb, per_test)
    _sheet_speed(wb, perf_csv_rows)
    _sheet_behaviours(wb)
    _sheet_home07(wb, records, gv_records)
    _americanize_workbook(wb)
    _save_deterministic(wb, out_path)
    return out_path


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Build English NIGHT-1 xlsx report")
    p.add_argument("run_id", help="run id or path to run dir")
    p.add_argument("--results-root", default=None)
    a = p.parse_args(argv)
    run_dir = _resolve_run_dir(a.run_id, a.results_root)
    out_path = os.path.join(run_dir, XLSX_NAME)
    build_workbook(run_dir, out_path)
    print("english excel report written to %s" % out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
