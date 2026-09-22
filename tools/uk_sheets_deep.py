"""Deep per-model sheets for the Ukrainian NIGHT-1 workbook.

Presentation layer only: renders the already-written Ukrainian per-model
analysis (``analysis_uk/group_{a,b,c,d}.json``) as five workbook sheets that
mirror the audited English sheets from ``tools/en_sheets_narrative.py`` and
``tools/en_sheets_detail.py`` (same structure, column layout, sorting,
tiering and formatting decisions), with Ukrainian sheet names, headers and
fixed strings.

Adds no claims about any model: every model-specific string is taken
verbatim from the input JSON. A missing field renders as ``Немає даних``.
Shared formatting helpers (``ACCENT``, ``_style_table``, ``_q_scale``) and
the quote-marker normalization are imported from
``tools/build_excel_report_en.py`` (read-only import) so the two workbooks
can never normalize evidence quotes differently. Never contacts Ollama.

Determinism: fixed column widths, fixed row-height rule, no timestamps,
no randomness; iteration order is always sorted or input order.
"""
from __future__ import annotations

import json
import os

SHEET_NAMES_UK = [
    "Оцінка моделей",
    "Рекомендації",
    "Сильні та слабкі сторони",
    "Деталі по тестах",
    "Аналіз помилок",
]

GROUP_FILES = ("group_a.json", "group_b.json", "group_c.json", "group_d.json")
EXPECTED_MODELS = 24

NOT_AVAILABLE_UK = "Немає даних"
BULLET = "\u2022"

SCORECARD_HEADERS_UK = [
    "Модель",
    "Оцінка (0–10)",
    "Тестів оцінено",
    "Покриття",
    "Підстава оцінки",
    "Що це за модель",
    "Роль у прогоні",
    "Підсумок",
]

RECOMMENDATION_HEADERS_UK = [
    "Модель",
    "Коли застосовувати",
    "Коли уникати",
    "Відповідність залізу (RTX 3070, 8 ГБ)",
]

STRENGTH_HEADERS_UK = [
    "Модель",
    "Сильні сторони",
    "Слабкі сторони",
]

DETAIL_HEADER_UK = ["Модель", "Тест", "Назва тесту", "Q_sem",
                    "Кейси (оцінено/всього)", "Результат", "Примітка"]
FAILURE_HEADER_UK = ["Модель", "Шаблон помилки", "Кількість", "Що це означає",
                     "Що зробила модель", "Приклад (кейс)", "Приклад виводу",
                     "Чому це важливо"]

TIER_GENERALIST_UK = "Моделі-універсали (5 і більше оцінених тестів)"
TIER_FOCUSED_UK = "Сфокусовані моделі (2–4 оцінені тести)"
TIER_SINGLE_UK = "Однотестові моделі (1 оцінений тест)"

COVERAGE_GENERALIST_UK = "Універсал (5+ тестів)"
COVERAGE_FOCUSED_UK = "Сфокусована (2–4 тести)"
COVERAGE_SINGLE_UK = "Один тест"

SCORE_NOTE_UK = (
    "Як читати оцінки: оцінки можна порівнювати лише в межах однієї групи, "
    "але не між групами. Модель, оцінену за одним вузьким тестом, і модель, "
    "оцінену за сімнадцятьма різними тестами, фактично питали про різне, "
    "тож вище число у вужчій групі не означає сильнішу модель загалом. "
    "Покриття наведено явно, щоб широту охоплення не плутали з місцем "
    "у рейтингу."
)

NO_PATTERN_ROW_UK = ("Жоден шаблон помилок не досяг порога звітування "
                     "в цьому прогоні")

# Fixed outcome vocabulary from the translation (no synonyms).
OUTCOME_UK = ("сильно", "задовільно", "слабко", "провал", "не виконувався")
NOT_ATTEMPTED_UK = "не виконувався"

# English aliases accepted on input only (the JSON already carries the
# Ukrainian vocabulary); output is always one of OUTCOME_UK.
_OUTCOME_ALIASES = {
    "strong": "сильно",
    "adequate": "задовільно",
    "weak": "слабко",
    "failed": "провал",
    "not attempted": "не виконувався",
}


def load_analysis_uk(analysis_dir: str) -> list[dict]:
    """Read the four Ukrainian group files, merge, sort by model tag.

    Raises ValueError when a group file is missing, when a model tag is
    duplicated across groups, or when the merged result is not exactly
    24 unique models.
    """
    merged: list[dict] = []
    for name in GROUP_FILES:
        path = os.path.join(analysis_dir, name)
        if not os.path.isfile(path):
            raise ValueError("missing group file: %s" % path)
        with open(path, encoding="utf-8") as f:
            group = json.load(f)
        if not isinstance(group, list):
            raise ValueError("group file is not a list: %s" % path)
        merged.extend(group)
    tags = [m.get("model") for m in merged]
    if any(not isinstance(t, str) or not t for t in tags):
        raise ValueError("every entry must carry a non-empty 'model' tag")
    seen: set[str] = set()
    dupes = sorted({t for t in tags if t in seen or seen.add(t)})
    if dupes:
        raise ValueError("duplicated model tag(s): %s" % ", ".join(dupes))
    if len(merged) != EXPECTED_MODELS:
        raise ValueError(
            "expected %d unique models, found %d in %s"
            % (EXPECTED_MODELS, len(merged), analysis_dir)
        )
    merged.sort(key=lambda m: str(m.get("model")))
    return merged


def uk_count(n: int, one: str, few: str, many: str) -> str:
    """Render ``n`` with a numeral-agreeing Ukrainian noun.

    ``one``/``few``/``many`` are the 1/2-4/5+ forms, e.g. ("раз", "рази",
    "разів"). The 11-14 exception always takes ``many``.
    """
    m = abs(int(n)) % 100
    if 11 <= m <= 14:
        form = many
    elif m % 10 == 1:
        form = one
    elif 2 <= m % 10 <= 4:
        form = few
    else:
        form = many
    return "%d %s" % (int(n), form)


def uk_raziv(n: int) -> str:
    """``n`` with раз/рази/разів agreement."""
    return uk_count(n, "раз", "рази", "разів")


def uk_zapisiv(n: int) -> str:
    """``n`` with запис/записи/записів agreement."""
    return uk_count(n, "запис", "записи", "записів")


def uk_vipadkiv(n: int) -> str:
    """``n`` with випадок/випадки/випадків agreement."""
    return uk_count(n, "випадок", "випадки", "випадків")


def _tier_of(tests_scored: object) -> tuple[str, str]:
    """Return (band label, coverage cell) for a tests_scored count."""
    try:
        n = int(tests_scored)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        n = 0
    if n >= 5:
        return TIER_GENERALIST_UK, COVERAGE_GENERALIST_UK
    if n >= 2:
        return TIER_FOCUSED_UK, COVERAGE_FOCUSED_UK
    return TIER_SINGLE_UK, COVERAGE_SINGLE_UK


def _text(value: object) -> str:
    if value is None:
        return NOT_AVAILABLE_UK
    s = str(value).strip()
    return s if s else NOT_AVAILABLE_UK


def _outcome_label(raw: object) -> str:
    key = str(raw or "").strip().lower()
    if key in OUTCOME_UK:
        return key
    if key in _OUTCOME_ALIASES:
        return _OUTCOME_ALIASES[key]
    return NOT_AVAILABLE_UK


def _is_none_observed(items: list) -> bool:
    """True for a single honest 'no ... observed' sentence (either language).

    Only the one-item no-evidence sentence renders without a bullet; any
    other single-item list is real content and keeps its bullet.
    """
    import re as _re

    return (
        len(items) == 1
        and isinstance(items[0], str)
        and bool(_re.match(r"\s*(no\b.*\bobserv|немає\b|не виявлено\b|"
                           r"відсутн)", items[0], _re.IGNORECASE))
    )


def _bullet_block(items: object) -> str:
    """Render a list as one bulleted cell block, or a fallback string."""
    if not isinstance(items, list) or not items:
        return NOT_AVAILABLE_UK
    cleaned = [str(x).strip() for x in items if str(x).strip()]
    if not cleaned:
        return NOT_AVAILABLE_UK
    if _is_none_observed(cleaned):
        return cleaned[0]
    return "\n".join("%s %s" % (BULLET, s) for s in cleaned)


def normalize_quote_uk(quote: object, full_content: object) -> tuple[str, str]:
    """Normalize one evidence quote against its record's full content.

    Thin delegation to the shared English-builder helper
    (``tools/build_excel_report_en.py::normalize_evidence_quote``), so both
    workbooks normalize quotes identically. Returns (rendered, status);
    the quote text itself is never altered, only the trailing marker.
    """
    from tools.build_excel_report_en import normalize_evidence_quote

    return normalize_evidence_quote(quote, full_content)


def strip_quote_marker_uk(quote: str) -> str:
    """Remove one trailing truncation marker from an evidence quote.

    Thin delegation to the shared English-builder helper; handles the
    markers observed across the four analysis groups (``[truncated]``,
    ``[...]``) plus a bare trailing ellipsis.
    """
    from tools.build_excel_report_en import strip_trailing_marker

    return strip_trailing_marker(quote)


def _header_fill():
    from openpyxl.styles import PatternFill

    from tools.build_excel_report_en import ACCENT

    return PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")


def _paint_header_row(ws, row_idx: int, n_cols: int) -> None:
    from openpyxl.styles import Alignment, Font

    fill = _header_fill()
    font = Font(bold=True, color="FFFFFF")
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _rating_scale(ws, col_letter: str, first: int, last: int) -> None:
    """Navy color scale over 0-10 ratings, applied to one tier's rows."""
    try:
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError:
        return
    if last < first:
        return
    from tools.build_excel_report_en import ACCENT

    ws.conditional_formatting.add(
        "%s%d:%s%d" % (col_letter, first, col_letter, last),
        ColorScaleRule(
            start_type="num", start_value=0.0, start_color="F2F2F2",
            mid_type="num", mid_value=5.0, mid_color="8EA9DB",
            end_type="num", end_value=10.0, end_color=ACCENT,
        ),
    )


def _rating_of(model: dict) -> float | None:
    overall = model.get("overall") or {}
    r = overall.get("rating_10") if isinstance(overall, dict) else None
    if isinstance(r, bool) or not isinstance(r, (int, float)):
        return None
    return float(r)


def build_scorecard_sheet(wb, analysis: list[dict]) -> None:
    """Sheet 1 -- one row per model, grouped in three coverage tiers."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet(SHEET_NAMES_UK[0])
    n_cols = len(SCORECARD_HEADERS_UK)
    last_col = ws.cell(row=1, column=n_cols).column_letter

    ws.merge_cells("A1:%s1" % last_col)
    title = ws["A1"]
    title.value = "Оцінка моделей"
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(vertical="center")

    ws.merge_cells("A2:%s2" % last_col)
    note = ws["A2"]
    note.value = SCORE_NOTE_UK
    note.font = Font(italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 60

    for col, header in enumerate(SCORECARD_HEADERS_UK, 1):
        ws.cell(row=3, column=col, value=header)
    _paint_header_row(ws, 3, n_cols)

    tiers: list[tuple[str, list[dict]]] = [
        (TIER_GENERALIST_UK, []),
        (TIER_FOCUSED_UK, []),
        (TIER_SINGLE_UK, []),
    ]
    band_of = {
        TIER_GENERALIST_UK: tiers[0],
        TIER_FOCUSED_UK: tiers[1],
        TIER_SINGLE_UK: tiers[2],
    }
    for m in analysis:
        overall = m.get("overall") if isinstance(m.get("overall"), dict) else {}
        band, _ = _tier_of(overall.get("tests_scored"))
        band_of[band][1].append(m)
    for _, members in tiers:
        members.sort(
            key=lambda m: (
                -_rating_of(m) if _rating_of(m) is not None else float("inf"),
                str(m.get("model")),
            )
        )

    band_fill = PatternFill(
        start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
    )
    row = 4
    for band_label, members in tiers:
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=n_cols)
        band = ws.cell(row=row, column=1, value=band_label)
        band.font = Font(bold=True)
        band.fill = band_fill
        band.alignment = Alignment(vertical="center")
        row += 1
        tier_first = row
        for m in members:
            overall = m.get("overall") if isinstance(
                m.get("overall"), dict) else {}
            profile = m.get("profile") if isinstance(
                m.get("profile"), dict) else {}
            _, coverage = _tier_of(overall.get("tests_scored"))
            rating = _rating_of(m)
            ws.cell(row=row, column=1, value=str(m.get("model")))
            ws.cell(row=row, column=2,
                    value=rating if rating is not None else NOT_AVAILABLE_UK)
            n_tests = overall.get("tests_scored")
            ws.cell(row=row, column=3,
                    value=n_tests if isinstance(n_tests, int)
                    and not isinstance(n_tests, bool) else NOT_AVAILABLE_UK)
            ws.cell(row=row, column=4, value=coverage)
            ws.cell(row=row, column=5, value=_text(overall.get("rating_basis")))
            ws.cell(row=row, column=6, value=_text(profile.get("documented")))
            ws.cell(row=row, column=7, value=_text(profile.get("role_in_run")))
            ws.cell(row=row, column=8, value=_text(m.get("bottom_line")))
            for col in range(5, 9):
                ws.cell(row=row, column=col).alignment = Alignment(
                    wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 90
            row += 1
        _rating_scale(ws, "B", tier_first, row - 1)

    ws.freeze_panes = "A4"
    widths = [24, 12, 12, 18, 45, 55, 55, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[
            ws.cell(row=3, column=i).column_letter].width = w
    for r in ws.iter_rows(min_row=4):
        if isinstance(r[1].value, (int, float)):
            r[1].number_format = "0.0"
        if isinstance(r[2].value, int):
            r[2].number_format = "0"


def _bulleted_sheet(wb, title: str, headers: list[str], widths: list[float],
                    rows: list[list[str]]) -> None:
    from openpyxl.styles import Alignment

    from tools.build_excel_report_en import _style_table

    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    _style_table(ws, widths)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    for ws_row in ws.iter_rows(min_row=2, max_col=len(headers)):
        longest = 1
        for cell in ws_row[1:]:
            cell.alignment = wrap_top
            if isinstance(cell.value, str):
                longest = max(longest, cell.value.count("\n") + 1)
        ws.row_dimensions[ws_row[0].row].height = min(
            220, 30 + 22 * longest)


def build_recommendations_sheet(wb, analysis: list[dict]) -> None:
    """Sheet 2 -- use-when / avoid-when / hardware fit, one row per model."""
    rows = []
    for m in sorted(analysis, key=lambda x: str(x.get("model"))):
        rows.append([
            str(m.get("model")),
            _bullet_block(m.get("use_when")),
            _bullet_block(m.get("avoid_when")),
            _text(m.get("hardware_fit")),
        ])
    _bulleted_sheet(wb, SHEET_NAMES_UK[1], RECOMMENDATION_HEADERS_UK,
                    [24, 60, 60, 65], rows)


def build_strengths_sheet(wb, analysis: list[dict]) -> None:
    """Sheet 3 -- strengths and weaknesses, one row per model."""
    rows = []
    for m in sorted(analysis, key=lambda x: str(x.get("model"))):
        rows.append([
            str(m.get("model")),
            _bullet_block(m.get("strengths")),
            _bullet_block(m.get("weaknesses")),
        ])
    _bulleted_sheet(wb, SHEET_NAMES_UK[2], STRENGTH_HEADERS_UK,
                    [24, 65, 65], rows)


def build_all_narrative_uk(wb, analysis: list[dict]) -> None:
    """Build all three Ukrainian narrative sheets into an open workbook."""
    build_scorecard_sheet(wb, analysis)
    build_recommendations_sheet(wb, analysis)
    build_strengths_sheet(wb, analysis)


def build_detail_sheet(wb, analysis: list[dict]):
    """Append the 'Деталі по тестах' sheet (one row per model x test)."""
    from openpyxl.styles import Alignment, Border, Side
    from tools.build_excel_report_en import _q_scale, _style_table

    ws = wb.create_sheet(SHEET_NAMES_UK[3])
    ws.append(DETAIL_HEADER_UK)
    first_rows: set[int] = set()
    for model in analysis:
        tag = str(model.get("model", "") or "").strip() or NOT_AVAILABLE_UK
        entries = sorted(model.get("per_test") or [],
                         key=lambda e: str(e.get("test_id", "")))
        first_rows.add(ws.max_row + 1)
        for entry in entries:
            outcome_raw = str(entry.get("outcome", "") or "").strip().lower()
            not_attempted = outcome_raw == NOT_ATTEMPTED_UK
            q_raw = entry.get("q_sem")
            if not_attempted:
                q_cell: object = None  # genuinely empty: no score exists
            elif isinstance(q_raw, bool) or q_raw is None or (
                    isinstance(q_raw, str) and not q_raw.strip()):
                q_cell = NOT_AVAILABLE_UK
            else:
                try:
                    q_cell = round(float(q_raw), 3)
                except (TypeError, ValueError):
                    q_cell = NOT_AVAILABLE_UK
            ws.append([tag,
                       _text(entry.get("test_id")),
                       _text(entry.get("title")),
                       q_cell,
                       None if not_attempted else _text(entry.get("n")),
                       _outcome_label(entry.get("outcome")),
                       _text(entry.get("note"))])
    _style_table(ws, [22, 10, 40, 10, 18, 14, 72], {3: "0.000"})
    ws.freeze_panes = "B2"  # keep the header row and the Model column visible
    ws.auto_filter.ref = ws.dimensions
    _q_scale(ws, "D", ws.max_row)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    block_top = Border(top=Side(style="thin", color="BFBFBF"))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                            max_col=len(DETAIL_HEADER_UK)):
        for cell in row:
            cell.alignment = wrap_top
            if cell.row in first_rows:
                cell.border = block_top
        ws.row_dimensions[row[0].row].height = 30
    return ws


def build_failures_sheet(wb, analysis: list[dict]):
    """Append the 'Аналіз помилок' sheet (one row per pattern per model).

    Every model gets at least one row: models with no qualifying patterns
    get a single explanatory row, so filtering by model never returns an
    empty result. Evidence quotes render byte-identical to the analysis
    JSON (no translation, no reformatting); any truncation-marker
    normalization happens upstream through the shared helper
    (see ``normalize_quote_uk``), never here.
    """
    from openpyxl.styles import Alignment, Font
    from tools.build_excel_report_en import _style_table

    ws = wb.create_sheet(SHEET_NAMES_UK[4])
    ws.append(FAILURE_HEADER_UK)
    quote_font = Font(name="Consolas")
    for model in analysis:
        tag = str(model.get("model", "") or "").strip() or NOT_AVAILABLE_UK
        patterns = sorted(model.get("failure_analysis") or [],
                          key=lambda e: (-(e.get("count")
                                           if isinstance(e.get("count"), int)
                                           and not isinstance(e.get("count"), bool)
                                           else -1),
                                         str(e.get("pattern", ""))))
        if not patterns:
            ws.append([tag, NO_PATTERN_ROW_UK, None, None,
                       None, None, None, None])
            continue
        for entry in patterns:
            count = entry.get("count")
            count_cell: object = (count if isinstance(count, int)
                                  and not isinstance(count, bool)
                                  else NOT_AVAILABLE_UK)
            ws.append([tag,
                       _text(entry.get("pattern")),
                       count_cell,
                       _text(entry.get("english_meaning")),
                       _text(entry.get("what_the_model_did")),
                       _text(entry.get("example_case_id")),
                       entry.get("example_quote")
                       if isinstance(entry.get("example_quote"), str)
                       and entry.get("example_quote", "").strip()
                       else NOT_AVAILABLE_UK,
                       _text(entry.get("why_it_matters"))])
    _style_table(ws, [22, 30, 8, 42, 52, 14, 52, 42])
    ws.auto_filter.ref = ws.dimensions
    wrap_top = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                            max_col=len(FAILURE_HEADER_UK)):
        for cell in row:
            cell.alignment = wrap_top
        ws.row_dimensions[row[0].row].height = 60
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7,
                            max_row=ws.max_row):
        for cell in row:
            if isinstance(cell.value, str) and cell.value != NOT_AVAILABLE_UK:
                cell.font = quote_font  # verbatim model output: monospace
    return ws


def build_all_uk(wb, analysis: list[dict]) -> None:
    """Build all five Ukrainian deep sheets into an open workbook."""
    build_all_narrative_uk(wb, analysis)
    build_detail_sheet(wb, analysis)
    build_failures_sheet(wb, analysis)
