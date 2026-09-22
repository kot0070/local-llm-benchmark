"""Per-model narrative sheets for the English NIGHT-1 workbook.

Presentation layer only: renders the already-written per-model analysis
(``analysis_en/group_{a,b,c,d}.json``, see ``tasks/DEEP_ANALYSIS_SPEC.md``)
as three workbook sheets that answer the owner's three questions -- what
each model is and how it rated, when to use or avoid it, and what its
strengths and weaknesses are.

Adds no claims about any model: every model-specific string is taken
verbatim from the input JSON. A missing field renders as ``Not available``.
All wording added by this module (headers, tier labels, notes) is plain
US English. Stdlib + openpyxl only. Never contacts Ollama.

Determinism: fixed column widths, fixed row-height rule, no timestamps,
no randomness; iteration order is always sorted or input order.
"""
from __future__ import annotations

import json
import os

SHEET_NAMES_NARRATIVE = [
    "Model Scorecard",
    "Recommendations",
    "Strengths & Weaknesses",
]

GROUP_FILES = ("group_a.json", "group_b.json", "group_c.json", "group_d.json")
EXPECTED_MODELS = 24

NOT_AVAILABLE = "Not available"
BULLET = "\u2022"

SCORECARD_HEADERS = [
    "Model",
    "Rating (0-10)",
    "Tests scored",
    "Coverage",
    "Rating basis",
    "What it is",
    "Role in this run",
    "Bottom line",
]

RECOMMENDATION_HEADERS = [
    "Model",
    "Use when",
    "Avoid when",
    "Hardware fit (RTX 3070, 8 GB)",
]

STRENGTH_HEADERS = [
    "Model",
    "Strengths",
    "Weaknesses",
]

TIER_GENERALIST = "Generalist models (5 or more tests scored)"
TIER_FOCUSED = "Focused models (2 to 4 tests scored)"
TIER_SINGLE = "Single-test models (1 test scored)"

COVERAGE_GENERALIST = "Generalist (5+ tests)"
COVERAGE_FOCUSED = "Focused (2-4 tests)"
COVERAGE_SINGLE = "Single-test (1 test)"

SCORE_NOTE = (
    "How to read the ratings: ratings can be compared within a tier, but not "
    "across tiers. A model judged on a single narrow task and a model judged "
    "on seventeen varied tasks were not asked the same question, so a higher "
    "number in a narrower tier does not mean a stronger model overall. "
    "Coverage is shown explicitly so breadth is never mistaken for rank."
)


def load_analysis(analysis_dir: str) -> list[dict]:
    """Read the four group files, merge, and sort by model tag.

    Raises ValueError when a group file is missing, when a model tag is
    duplicated across groups, or when the merged result is not exactly
    24 unique models.
    """
    merged: list[dict] = []
    for name in GROUP_FILES:
        path = os.path.join(analysis_dir, name)
        if not os.path.isfile(path):
            raise ValueError("missing group file: %s" % path)
        with open(path, encoding="utf-8") as f:
            group = json.load(f)
        if not isinstance(group, list):
            raise ValueError("group file is not a list: %s" % path)
        merged.extend(group)
    tags = [m.get("model") for m in merged]
    if any(not isinstance(t, str) or not t for t in tags):
        raise ValueError("every entry must carry a non-empty 'model' tag")
    seen: set[str] = set()
    dupes = sorted({t for t in tags if t in seen or seen.add(t)})
    if dupes:
        raise ValueError("duplicated model tag(s): %s" % ", ".join(dupes))
    if len(merged) != EXPECTED_MODELS:
        raise ValueError(
            "expected %d unique models, found %d in %s"
            % (EXPECTED_MODELS, len(merged), analysis_dir)
        )
    merged.sort(key=lambda m: str(m.get("model")))
    return merged


def _tier_of(tests_scored: object) -> tuple[str, str]:
    """Return (band label, coverage cell) for a tests_scored count."""
    try:
        n = int(tests_scored)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        n = 0
    if n >= 5:
        return TIER_GENERALIST, COVERAGE_GENERALIST
    if n >= 2:
        return TIER_FOCUSED, COVERAGE_FOCUSED
    return TIER_SINGLE, COVERAGE_SINGLE


def _text(value: object) -> str:
    if value is None:
        return NOT_AVAILABLE
    s = str(value).strip()
    return s if s else NOT_AVAILABLE


def _is_none_observed(items: list) -> bool:
    """True for a single honest 'no ... observed' sentence.

    Only the one-item no-evidence sentence renders without a bullet; any
    other single-item list is real content and keeps its bullet.
    """
    import re as _re

    return (
        len(items) == 1
        and isinstance(items[0], str)
        and bool(_re.match(r"\s*no\b.*\bobserv", items[0], _re.IGNORECASE))
    )


def _bullet_block(items: object) -> str:
    """Render a list as one bulleted cell block, or a fallback string."""
    if not isinstance(items, list) or not items:
        return NOT_AVAILABLE
    cleaned = [str(x).strip() for x in items if str(x).strip()]
    if not cleaned:
        return NOT_AVAILABLE
    if _is_none_observed(cleaned):
        return cleaned[0]
    return "\n".join("%s %s" % (BULLET, s) for s in cleaned)


def _header_fill():
    from openpyxl.styles import PatternFill

    from tools.build_excel_report_en import ACCENT

    return PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")


def _paint_header_row(ws, row_idx: int, n_cols: int) -> None:
    from openpyxl.styles import Alignment, Font

    fill = _header_fill()
    font = Font(bold=True, color="FFFFFF")
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _rating_scale(ws, col_letter: str, first: int, last: int) -> None:
    """Navy color scale over 0-10 ratings, applied to one tier's rows."""
    try:
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError:
        return
    if last < first:
        return
    from tools.build_excel_report_en import ACCENT

    ws.conditional_formatting.add(
        "%s%d:%s%d" % (col_letter, first, col_letter, last),
        ColorScaleRule(
            start_type="num", start_value=0.0, start_color="F2F2F2",
            mid_type="num", mid_value=5.0, mid_color="8EA9DB",
            end_type="num", end_value=10.0, end_color=ACCENT,
        ),
    )


def _rating_of(model: dict) -> float | None:
    overall = model.get("overall") or {}
    r = overall.get("rating_10") if isinstance(overall, dict) else None
    if isinstance(r, bool) or not isinstance(r, (int, float)):
        return None
    return float(r)


def build_scorecard_sheet(wb, analysis: list[dict]) -> None:
    """Sheet 1 -- one row per model, grouped in three coverage tiers."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet(SHEET_NAMES_NARRATIVE[0])
    n_cols = len(SCORECARD_HEADERS)
    last_col = ws.cell(row=1, column=n_cols).column_letter

    ws.merge_cells("A1:%s1" % last_col)
    title = ws["A1"]
    title.value = "Model scorecard"
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(vertical="center")

    ws.merge_cells("A2:%s2" % last_col)
    note = ws["A2"]
    note.value = SCORE_NOTE
    note.font = Font(italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 60

    for col, header in enumerate(SCORECARD_HEADERS, 1):
        ws.cell(row=3, column=col, value=header)
    _paint_header_row(ws, 3, n_cols)

    tiers: list[tuple[str, list[dict]]] = [
        (TIER_GENERALIST, []),
        (TIER_FOCUSED, []),
        (TIER_SINGLE, []),
    ]
    band_of = {
        TIER_GENERALIST: tiers[0],
        TIER_FOCUSED: tiers[1],
        TIER_SINGLE: tiers[2],
    }
    for m in analysis:
        overall = m.get("overall") if isinstance(m.get("overall"), dict) else {}
        band, _ = _tier_of(overall.get("tests_scored"))
        band_of[band][1].append(m)
    for _, members in tiers:
        members.sort(
            key=lambda m: (
                -_rating_of(m) if _rating_of(m) is not None else float("inf"),
                str(m.get("model")),
            )
        )

    band_fill = PatternFill(
        start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
    )
    row = 4
    for band_label, members in tiers:
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=n_cols)
        band = ws.cell(row=row, column=1, value=band_label)
        band.font = Font(bold=True)
        band.fill = band_fill
        band.alignment = Alignment(vertical="center")
        row += 1
        tier_first = row
        for m in members:
            overall = m.get("overall") if isinstance(
                m.get("overall"), dict) else {}
            profile = m.get("profile") if isinstance(
                m.get("profile"), dict) else {}
            _, coverage = _tier_of(overall.get("tests_scored"))
            rating = _rating_of(m)
            ws.cell(row=row, column=1, value=str(m.get("model")))
            ws.cell(row=row, column=2,
                    value=rating if rating is not None else NOT_AVAILABLE)
            n_tests = overall.get("tests_scored")
            ws.cell(row=row, column=3,
                    value=n_tests if isinstance(n_tests, int)
                    and not isinstance(n_tests, bool) else NOT_AVAILABLE)
            ws.cell(row=row, column=4, value=coverage)
            ws.cell(row=row, column=5, value=_text(overall.get("rating_basis")))
            ws.cell(row=row, column=6, value=_text(profile.get("documented")))
            ws.cell(row=row, column=7, value=_text(profile.get("role_in_run")))
            ws.cell(row=row, column=8, value=_text(m.get("bottom_line")))
            for col in range(5, 9):
                ws.cell(row=row, column=col).alignment = Alignment(
                    wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 90
            row += 1
        _rating_scale(ws, "B", tier_first, row - 1)

    ws.freeze_panes = "A4"
    widths = [24, 12, 12, 18, 45, 55, 55, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[
            ws.cell(row=3, column=i).column_letter].width = w
    for r in ws.iter_rows(min_row=4):
        if isinstance(r[1].value, (int, float)):
            r[1].number_format = "0.0"
        if isinstance(r[2].value, int):
            r[2].number_format = "0"


def _bulleted_sheet(wb, title: str, headers: list[str], widths: list[float],
                    rows: list[list[str]]) -> None:
    from openpyxl.styles import Alignment

    from tools.build_excel_report_en import _style_table

    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    _style_table(ws, widths)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    for ws_row in ws.iter_rows(min_row=2, max_col=len(headers)):
        longest = 1
        for cell in ws_row[1:]:
            cell.alignment = wrap_top
            if isinstance(cell.value, str):
                longest = max(longest, cell.value.count("\n") + 1)
        ws.row_dimensions[ws_row[0].row].height = min(
            220, 30 + 22 * longest)


def build_recommendations_sheet(wb, analysis: list[dict]) -> None:
    """Sheet 2 -- use-when / avoid-when / hardware fit, one row per model."""
    rows = []
    for m in sorted(analysis, key=lambda x: str(x.get("model"))):
        rows.append([
            str(m.get("model")),
            _bullet_block(m.get("use_when")),
            _bullet_block(m.get("avoid_when")),
            _text(m.get("hardware_fit")),
        ])
    _bulleted_sheet(wb, SHEET_NAMES_NARRATIVE[1], RECOMMENDATION_HEADERS,
                    [24, 60, 60, 65], rows)


def build_strengths_sheet(wb, analysis: list[dict]) -> None:
    """Sheet 3 -- strengths and weaknesses, one row per model."""
    rows = []
    for m in sorted(analysis, key=lambda x: str(x.get("model"))):
        rows.append([
            str(m.get("model")),
            _bullet_block(m.get("strengths")),
            _bullet_block(m.get("weaknesses")),
        ])
    _bulleted_sheet(wb, SHEET_NAMES_NARRATIVE[2], STRENGTH_HEADERS,
                    [24, 65, 65], rows)


def build_all_narrative(wb, analysis: list[dict]) -> None:
    """Build all three narrative sheets into an open workbook."""
    build_scorecard_sheet(wb, analysis)
    build_recommendations_sheet(wb, analysis)
    build_strengths_sheet(wb, analysis)
